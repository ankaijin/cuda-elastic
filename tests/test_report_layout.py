import importlib.util
import os
from pathlib import Path
import tempfile
import types
import unittest
from unittest import mock

from openpyxl import Workbook, load_workbook


class ReportLayoutTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        path = Path(__file__).resolve().parents[1] / 'results/build_spreadsheet.py'
        source = Path(os.environ.get('REPORT_BUILDER_SOURCE', str(path))).read_text()
        cls.m = types.ModuleType('layout_builder')
        cls.m.__file__ = str(path)
        exec(compile(source, str(path), 'exec'), cls.m.__dict__)

    def test_summary_preserves_application_values_and_speedups(self):
        m = self.m
        row = ('mlperf-tiny VWW', 'Whole application', '16 bytes', m.TIMED,
               m.Cell(12.0), m.Cell(8.0), m.Cell(4.0), 'A measured example.')
        wb = Workbook()
        wb.remove(wb.active)
        sheet = m.build_summary(wb, m.RTX3090, [row])
        self.assertEqual([sheet.cell(5, c).value for c in range(5, 10)],
                         [12.0, 8.0, 4.0, 1.5, 2.0])
        self.assertEqual(sheet['J5'].comment.text, 'A measured example.')

    def test_5090_supporting_sheets_keep_missing_measurements_pending(self):
        m = self.m
        gpu = m.RTX5090
        rows = m.gpu_rows(gpu)
        wb = Workbook()
        wb.remove(wb.active)
        summary = m.build_summary(wb, gpu, rows, base_rows=rows)
        m.build_latency(wb, gpu, rows)
        m.build_summary(wb, gpu, rows, base_rows=rows, kernel=True)
        m.build_sources(wb, [(gpu, rows)])
        self.assertEqual(summary.max_row, 10)
        for row in range(5, summary.max_row + 1):
            self.assertTrue(all(summary.cell(row, col).value == 'pending' for col in range(5, 8)))
            self.assertEqual(summary.cell(row, 10).value, 'pending')
        latency = wb['Latency RTX 5090']
        self.assertEqual(latency.max_row, 22)
        self.assertTrue(all('small' not in latency.cell(row, 2).value
                            for row in range(5, latency.max_row + 1)))
        self.assertTrue({'RTX 5090', 'Latency RTX 5090', 'Kernels RTX 5090', 'Sources'} <= set(wb.sheetnames))

    def test_generated_reports_keep_large_results_primary_and_small_results_diagnostic(self):
        m = self.m
        stages = [
            ('New diagnostic', 'Whole application, small shape (launch-bound)', False),
            ('New workload A', 'Whole application', False),
            ('New workload A', 'Kernel', True),
            ('New workload B', 'Whole application', False),
        ]
        # Put the diagnostic first and use distinct timings to catch baseline or
        # prediction misalignment after filtering, including for future workloads.
        timings = ([900, 100, 40, 200], [300, 25, 10, 50], [150, 20, 8, 40])
        fixtures = {}
        for gpu, values in zip(m.GPUS, timings):
            fixtures[gpu] = [
                (name, stage, 'fixture size', m.TIMED,
                 m.Cell(2 * value, kernel=kernel), m.Cell(value, kernel=kernel),
                 m.Cell(value / 2, kernel=kernel), '')
                for (name, stage, kernel), value in zip(stages, values)
            ]
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            with mock.patch.multiple(
                m, REPO=root, OUT=root / 'report.xlsx', MD_OUT=root / 'README.md',
                validate_inputs=lambda gpus: None, gpu_rows=fixtures.__getitem__,
                predicted_speedup=lambda cell, gpu: (cell.us, 'Fixture prediction'),
                link=lambda path: 'https://example.test/fixture.log',
            ):
                m.main()
            wb = load_workbook(root / 'report.xlsx')
            self.assertEqual(wb.sheetnames[:3], [gpu.sheet for gpu in m.GPUS])
            self.assertEqual([name for name in wb.sheetnames if name.startswith('Small ')],
                             ['Small RTX 3090'])
            for gpu, values in zip(m.GPUS, timings):
                summary = list(wb[gpu.sheet].iter_rows(min_row=5, values_only=True))
                self.assertEqual([(r[0], r[5]) for r in summary],
                                 [('New workload A', values[1]), ('New workload B', values[3])])
                if gpu is not m.BASELINE:
                    expected_speedup = 100 / values[1]
                    self.assertEqual([r[9] for r in summary], [expected_speedup] * 2)
                latency = list(wb[f'Latency {gpu.sheet}'].iter_rows(min_row=5, values_only=True))
                expected_names = {'New workload A', 'New workload B'}
                if gpu is m.BASELINE:
                    expected_names.add('New diagnostic')
                self.assertEqual({r[0] for r in latency}, expected_names)
                self.assertEqual(len(latency), 3 * len(expected_names))
            small = list(wb['Small RTX 3090'].iter_rows(min_row=5, values_only=True))
            self.assertEqual([(r[0], r[5]) for r in small], [('New diagnostic', 900)])
            analysis = list(wb['GPU analysis'].iter_rows(min_row=5, values_only=True))
            self.assertEqual([(r[0], r[1], r[4], r[5]) for r in analysis], [
                (gpu.sheet, name, prediction, speedup)
                for gpu, speedup in ((m.RTX4090, 4), (m.RTX5090, 5))
                for name, prediction in [('New workload A', 100), ('New workload A', 40),
                                         ('New workload B', 200)]
            ])
            for ws in wb:
                for row in ws:
                    for cell in row:
                        if cell.hyperlink and cell.hyperlink.target.startswith("#'"):
                            self.assertIn(cell.hyperlink.target.split("'", 2)[1], wb.sheetnames)
            headings = [line for line in (root / 'README.md').read_text().splitlines()
                        if line.startswith('## ')]
            self.assertEqual(headings, ['## RTX 3090', '## RTX 4090', '## RTX 5090',
                                        '## Small RTX 3090 diagnostics'])
            wb.close()

    def test_ncu_summary_keeps_configurations_and_reports_separate_and_preserves_launches(self):
        m = self.m

        def launch(duration, unit='us', grid=(2, 1, 1), dynamic='0', sm='20'):
            row = {
                'Kernel Name': ('example::compute(float *)', ''),
                'gpu__time_duration.sum': (duration, unit),
                'launch__grid_size': ('2', ''),
                'launch__block_size': ('256', ''),
                'launch__registers_per_thread': ('32', 'register/thread'),
                'launch__shared_mem_per_block': ('1', 'Kbyte/block'),
                'launch__shared_mem_per_block_dynamic': (dynamic, 'byte/block'),
                'sm__throughput.avg.pct_of_peak_sustained_elapsed': (sm, '%'),
                'gpu__dram_throughput.avg.pct_of_peak_sustained_elapsed': ('60', '%'),
                'sm__warps_active.avg.pct_of_peak_sustained_active': ('50', '%'),
            }
            for kind, dims in (('grid', grid), ('block', (256, 1, 1))):
                for axis, value in zip('xyz', dims):
                    row[f'launch__{kind}_dim_{axis}'] = (str(value), '')
            return row

        direct = [launch('1000', 'ns'), launch('3', sm='40'),
                  launch('5', grid=(1, 2, 1)), launch('7', dynamic='512')]
        del direct[1]['sm__warps_active.avg.pct_of_peak_sustained_active']
        reports = {'direct.ncu-rep': direct, 'agentic.ncu-rep': [launch('9')]}
        row = ('New workload', 'Kernel', 'large fixture', m.TIMED,
               m.Cell(ncu='direct.ncu-rep', kernel=True), m.Cell(kernel=True),
               m.Cell(ncu='agentic.ncu-rep', kernel=True), '')
        wb = Workbook()
        wb.remove(wb.active)
        with mock.patch.object(m, 'ncu_rows', side_effect=reports.__getitem__), \
                mock.patch.object(m, 'link', side_effect=lambda path: f'https://example.test/{path}'):
            m.build_ncu_details(wb, m.RTX3090, [row])
        summary, raw = wb['ncu metrics (RTX 3090)'], wb['NCU launches RTX 3090']
        self.assertEqual(summary.max_row, 8)
        self.assertEqual(raw.max_row, 9)
        self.assertEqual([summary.cell(r, 4).value for r in range(5, 9)], [2, 1, 1, 1])
        self.assertEqual([summary.cell(r, 5).value for r in range(5, 9)], [2, 5, 7, 9])
        self.assertEqual([summary.cell(5, c).value for c in (6, 7, 8, 9)], [1, 3, 30, 60])
        self.assertIsNone(summary.cell(5, 10).value)
        self.assertEqual(summary['C5'].value, 'compute')
        self.assertEqual(summary['C5'].comment.text, 'example::compute(float *)')
        self.assertEqual([summary.cell(r, 13).value for r in range(5, 8)],
                         ['2 x 1 x 1', '1 x 2 x 1', '2 x 1 x 1'])
        for r, version, start, count in ((5, 'Direct', 5, 2), (6, 'Direct', 7, 1),
                                         (7, 'Direct', 8, 1), (8, 'Agentic', 9, 1)):
            self.assertEqual(summary.cell(r, 18).hyperlink.target,
                             f"#'NCU launches RTX 3090'!A{start}")
            color = m.GROUP_FILLS[m.SHORT.index(version)]
            self.assertTrue(summary.cell(r, 1).fill.fgColor.rgb.endswith(color))
            self.assertEqual([raw.cell(rr, 2).value for rr in range(start, start + count)],
                             [version] * count)
        self.assertEqual([raw.cell(r, 4).value for r in range(5, 10)], [0, 1, 2, 3, 0])
        self.assertEqual(raw['F5'].value, '1000.0 ns')
        self.assertEqual(raw['F6'].value, '3.0 us')
        self.assertTrue(summary['Q8'].hyperlink.target.endswith('/agentic.ncu-rep'))
        self.assertEqual(summary.freeze_panes, 'D5')
        self.assertTrue(summary.auto_filter.ref)

    def test_5090_analysis_does_not_claim_a_hardware_prediction(self):
        m = self.m
        gpu = m.RTX5090
        rows = m.gpu_rows(gpu)
        wb = Workbook()
        wb.remove(wb.active)
        m.build_gpu_analysis(wb, [(m.BASELINE, rows), (gpu, rows)], {gpu: [None] * len(rows)})
        sheet = wb['GPU analysis']
        self.assertEqual(sheet.max_row, 16)
        for row in range(5, sheet.max_row + 1):
            self.assertNotIn('small', sheet.cell(row, 3).value)
            self.assertIsNone(sheet.cell(row, 5).value)
            self.assertEqual(sheet.cell(row, 6).value, 'pending')
            self.assertIn('not enabled', sheet.cell(row, 7).value.lower())


if __name__ == '__main__':
    unittest.main()
