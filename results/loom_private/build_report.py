#!/usr/bin/env python3
"""Build the loom_private v1 results page and workbook from selected records.

  python3 results/loom_private/build_report.py

Writes results/loom_private/README.md and results/loom_private/loom_private_v1.xlsx.
Inputs are pinned in results/loom_private/inputs/<gpu>.json, independently of
the legacy results/inputs manifests. Each application has these selections,
with paths relative to that GPU's folder (null means pending):

- `check` and `check_optimized`: `--action check` records of the baseline and
  the optimized variant.
- `bench_direct`, `bench_graph` (baseline) and `bench_optimized` (optimized
  variant, CUDA Graph launches): lists of `--action bench` records, one per
  recorded repetition. Times are the median over the list.
- `profiles`: `nsys` and `ncu` reports for `direct`, `graph` and `optimized`
  (the same three versions), written by applications/loom_private/tools/
  profile.py with CSV exports beside them, so reading them needs no Nsight
  tools.

Missing files, unsafe paths, records of another workload, variant, action or
launch (a check must cover both launch modes), records that validated an older
contract, records whose integer dimensions differ from the contract's `large`
values, records from a GPU other than the folder's, and profiles of another
application or version are errors.

Timing cells come only from bench records, so a correctness check never
appears as a performance measurement. The workbook uses the layout and styling
of results/cuda_performance.xlsx: per-GPU summary, Latency and Kernels sheets,
then Sources, GPU analysis, the ncu sheets, Correctness, Environment, and
Contracts. Its links point at GitHub, pinned to the latest origin/main commit
that has each file. The test-machine table comes from the recorder metadata
(<gpu>/runs/<run-id>/metadata.json) of the selected runs, except the source
revision, which the manifest states (`source_revision`, a full commit hash):
measurement runs start from commits that only add earlier runs' results, so
their recorded commits need not stay in the history.
Nothing is rerun. Requires openpyxl, and git for the workbook's links.
"""

import csv
import hashlib
import json
import re
import statistics
from pathlib import Path
import subprocess

REPO = Path(__file__).resolve().parents[2]
HERE = REPO / "results" / "loom_private"
SUITE = REPO / "applications" / "loom_private"
VERSION = "v1"
MD_OUT = HERE / "README.md"
XLSX_OUT = HERE / f"loom_private_{VERSION}.xlsx"
GPUS = ["rtx3090", "rtx4090", "rtx5090"]
APPS = ["pagerank", "attention", "decoder", "anomaly", "vecadd_memory", "vecadd_product"]
# Selection key -> (action, launch, variant) its records must have.
SELECTIONS = {"check": ("check", "both", "baseline"),
              "check_optimized": ("check", "both", "optimized"),
              "bench_direct": ("bench", "direct", "baseline"),
              "bench_graph": ("bench", "graph", "baseline"),
              "bench_optimized": ("bench", "graph", "optimized")}
CHECKS = ["check", "check_optimized"]
BENCHES = ["bench_direct", "bench_graph", "bench_optimized"]
# Profiled versions, in the order of BENCHES, and the file name each one ends with
# (tools/profile.py).
PROFILE_VERSIONS = {"direct": "baseline-direct", "graph": "baseline-graph",
                    "optimized": "optimized-graph"}
PROFILE_SUFFIXES = {"nsys": (".nsys-rep",), "ncu": (".ncu-rep", ".ncu-repz")}
# tools/profile.py runs one bench of --iterations 1 --warmup 0: four invocations.
PROFILED_INVOCATIONS = 4
PENDING = "pending"
# Display name and a format string over the contract's `large` values.
PROBLEMS = {
    "pagerank": ("PageRank", "{nodes:,} vertices, {edges:,} edges, {iterations} ordered "
                 "Gauss–Seidel iterations"),
    "attention": ("Multisensor attention", "{query_tokens:,} queries × {context_tokens:,} "
                  "context tokens"),
    "decoder": ("llama2c decoder", "{batch:,} sequences × {steps} steps, width {width:,}, "
                "vocabulary {vocab:,}"),
    "anomaly": ("Anomaly detection", "{samples:,} samples × {sample_elements} values through "
                "the {layers}-layer int8 model"),
    "vecadd_memory": ("Vecadd memory", "{length:,} integer additions"),
    "vecadd_product": ("Vecadd product", "{records:,} pairs of {operand_bytes}-byte operands, "
                       "added byte by byte"),
}
# What each optimized variant changes, in one sentence; its README section has
# the design and the reasons for each GPU's parameters.
OPTIMIZATIONS = {
    "pagerank": "One persistent cooperative kernel runs all sweeps. Per-node sweep tags replace "
                "the per-level launches, and the ordered Gauss–Seidel updates stay bit for bit "
                "the same, over a level-ordered layout of the same graph.",
    "attention": "Blocks of 64 queries share key tiles staged in shared memory, which removes "
                 "the L2 bottleneck behind the baseline's slow RTX 5090 time. A bound replaces "
                 "the row-maximum pass, and the context upload overlaps the computation.",
    "decoder": "A register-tiled classifier with a stream-K tail keeps each logit's sequential "
               "sum, so its output matches the baseline bit for bit.",
    "anomaly": "Int8 tensor-core tiles of 16 samples run all ten layers on chip. Requantization "
               "uses binary32 paths proven equal to the baseline's, and chunked uploads overlap "
               "inference and downloads.",
    "vecadd_memory": "Chunked uploads, kernels, and downloads on several streams overlap the two "
                     "copy directions, with chunk counts and vector widths chosen per GPU.",
    "vecadd_product": "The same chunked multi-stream pipeline around the baseline kernel, which "
                      "already runs near DRAM bandwidth.",
}
# Integer `large` entries that are not problem extents (left out of the size column).
NON_EXTENT_KEYS = {"expected_checksum"}
RUN_ID = re.compile(r"\d{8}-\d{6}-.+?-[0-9a-f]{8}(?=-)")


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def contract(app):
    path = SUITE / app / "contract.json"
    if not path.is_file():
        return None
    data = json.loads(path.read_text())
    data["_sha256"] = sha256(path)
    return data


def resolve(folder, value):
    """Exact path inside the GPU folder, or None when pending."""
    if value is None:
        return None
    if not isinstance(value, str) or not value or any(c in value for c in "*?["):
        raise ValueError(f"{folder}: selection must be an exact relative path: {value!r}")
    relative = Path(value)
    root = (REPO / folder).resolve()
    if relative.is_absolute() or ".." in relative.parts:
        raise ValueError(f"{folder}: selection escapes the GPU folder: {value}")
    path = (root / relative).resolve()
    if root not in path.parents:
        raise ValueError(f"{folder}: selection escapes the GPU folder: {value}")
    if not path.is_file():
        raise ValueError(f"{folder}: selected file is missing: {value}")
    return path


def is_integer(value):
    return isinstance(value, int) and not isinstance(value, bool)


def device_matches(folder, name):
    """True when the record's device name is the folder's GPU (rtx3090 -> RTX 3090)."""
    model = re.fullmatch(r"rtx(\d+)", folder)
    return (model is not None and isinstance(name, str)
            and re.fullmatch(rf"(NVIDIA )?(GeForce )?RTX {model.group(1)}", name) is not None)


def load_record(folder, app, key, value, spec):
    path = resolve(folder, value)
    if path is None:
        return None
    action, launch, variant = SELECTIONS[key]
    record = json.loads(path.read_text())
    where = f"{folder}/{value}"
    if record.get("schema") != "loom_private.result/1":
        raise ValueError(f"{where}: not a loom_private result record")
    if spec is None or record.get("identity") != spec.get("identity"):
        raise ValueError(f"{where}: record identity {record.get('identity')!r} "
                         f"is not application {app}")
    # Records from before variants existed are baselines.
    recorded = (record.get("action"), record.get("launch"), record.get("variant", "baseline"))
    if recorded != (action, launch, variant):
        raise ValueError(f"{where}: selected as {key} but records {'/'.join(recorded)}")
    if record.get("contract_sha256") != spec["_sha256"]:
        raise ValueError(f"{where}: record validated contract {record.get('contract_sha256')}, "
                         f"current contract is {spec['_sha256']}")
    if record.get("preset") != "large" or record.get("overrides"):
        raise ValueError(f"{where}: only unmodified large-preset records can be selected")
    dimensions = record.get("dimensions")
    if not isinstance(dimensions, dict):
        raise ValueError(f"{where}: record has no dimensions")
    for name, expected in spec.get("large", {}).items():
        actual = dimensions.get(name)
        if is_integer(expected) and is_integer(actual) and actual != expected:
            raise ValueError(f"{where}: dimension {name}={actual} differs from the "
                             f"contract's large value {expected}")
    device = (record.get("device") or {}).get("name")
    if not device_matches(folder, device):
        raise ValueError(f"{where}: recorded on {device!r}, not the {folder} folder's GPU")
    record["_path"] = f"{folder}/{value}"
    return record


def load_benches(folder, app, key, value, spec):
    """A bench selection: one path or a list of repetitions; None when pending."""
    if value is None:
        return None
    values = value if isinstance(value, list) else [value]
    if not values:
        raise ValueError(f"inputs/{folder}.json: {app} {key} is an empty list; use null")
    return [load_record(folder, app, key, v, spec) for v in values]


def load_profile(folder, app, kind, version, value):
    """A selected profile and the CSV export tools/profile.py wrote beside it."""
    path = resolve(folder, value)
    if path is None:
        return None
    suffix = next((s for s in PROFILE_SUFFIXES[kind] if path.name.endswith(s)), None)
    stem = path.name.removesuffix(suffix) if suffix else ""
    expected = PROFILE_VERSIONS[version]
    if (suffix is None or not stem.endswith(f"-{expected}")
            or Path(value).parts[:2] != ("loom_private", app)):
        raise ValueError(f"{folder}/{value}: not a {kind} report of {app} {expected}")
    export = path.with_name(f"{stem}_cuda_gpu_kern_sum.csv" if kind == "nsys"
                            else f"{stem}-ncu-raw.csv")
    if not export.is_file():
        raise ValueError(f"{folder}/{value}: its export {export.name} is missing")
    return {"_path": f"{folder}/{value}",
            "csv": f"{folder}/{Path(value).with_name(export.name)}"}


def load_gpu(folder):
    manifest = json.loads((HERE / "inputs" / f"{folder}.json").read_text())
    if manifest.get("folder") != folder or manifest.get("suite_version") != VERSION:
        raise ValueError(f"inputs/{folder}.json: folder or suite_version mismatch")
    if sorted(manifest["applications"]) != sorted(APPS):
        raise ValueError(f"inputs/{folder}.json: must list exactly {APPS}")
    rows = []
    for app in APPS:
        spec = contract(app)
        selected = manifest["applications"][app]
        if sorted(selected) != sorted([*SELECTIONS, "profiles"]):
            raise ValueError(f"inputs/{folder}.json: {app} needs keys "
                             f"{sorted([*SELECTIONS, 'profiles'])}")
        row = {"app": app, "contract": spec}
        for key in CHECKS:
            row[key] = load_record(folder, app, key, selected[key], spec)
        for key in BENCHES:
            row[key] = load_benches(folder, app, key, selected[key], spec)
        profiles = selected["profiles"]
        if sorted(profiles) != sorted(PROFILE_SUFFIXES) or any(
                sorted(profiles[kind]) != sorted(PROFILE_VERSIONS) for kind in profiles):
            raise ValueError(f"inputs/{folder}.json: {app} profiles need nsys and ncu, each "
                             f"with {sorted(PROFILE_VERSIONS)}")
        row["profiles"] = {(kind, version): load_profile(folder, app, kind, version,
                                                         profiles[kind][version])
                           for kind in PROFILE_SUFFIXES for version in PROFILE_VERSIONS}
        rows.append(row)
    return manifest, rows


def selected_records(row):
    """Every selected record of an application row."""
    found = [row[key] for key in CHECKS if row[key]]
    for key in BENCHES:
        found += row[key] or []
    return found


def timing(records, scope, field="batch_per_iteration_us"):
    """Median over the repetitions of a bench selection, else pending."""
    if not records:
        return PENDING
    return statistics.median(r["timing"][scope][field] for r in records)


def median_record(records, scope):
    """The repetition whose batched time is the (lower) median, for links."""
    ranked = sorted(records, key=lambda r: r["timing"][scope]["batch_per_iteration_us"])
    return ranked[(len(ranked) - 1) // 2]


def size_text(spec):
    if spec is None:
        return "contract pending"
    # Integer extents only; descriptive strings, nested notes, and non-extent
    # integers such as expected checksums stay in the contract.
    return ", ".join(f"{k}={v:,}" for k, v in spec.get("large", {}).items()
                     if is_integer(v) and k not in NON_EXTENT_KEYS and "checksum" not in k)


def check_cells(record):
    if record is None:
        return [PENDING] * 4
    def modes(launch):
        runs = [c for c in record["checks"] if c["launch"] == launch]
        return f"{sum(c['pass'] for c in runs)}/{len(runs)} pass" if runs else "not run"
    return [record["status"], modes("direct"), modes("graph"),
            "yes" if record["outputs_agree"] else "no"]


def fmt(value):
    if isinstance(value, str):
        return value
    return f"{value:,.0f}" if value >= 100 else f"{value:,.1f}"


def problem_text(app, spec):
    """One plain-language line describing the large problem, from the contract."""
    if spec is None:
        return "contract pending"
    return PROBLEMS[app][1].format(**spec["large"])


def check_summary(record):
    """pass only when every invocation passed and Direct and Graph agree."""
    if record is None:
        return PENDING
    passed = (record["status"] == "pass" and record["outputs_agree"]
              and all(c["pass"] for c in record["checks"]))
    return f"[{'pass' if passed else 'FAIL'}](../../{record['_path']})"


def run_metadata(path):
    """The recorder's metadata for the run that wrote a selected file, if any."""
    folder, _, name = path.partition("/")
    match = RUN_ID.match(Path(name).name)
    meta = REPO / folder / "runs" / match.group(0) / "metadata.json" if match else None
    return json.loads(meta.read_text()) if meta and meta.is_file() else None


def selected_runs(rows):
    """{run id or path: (kinds, metadata or None)} over every selected file."""
    runs = {}
    for row in rows:
        files = [(r["_path"], r["action"]) for r in selected_records(row)]
        files += [(p["_path"], "profile") for p in row["profiles"].values() if p]
        for path, kind in files:
            meta = run_metadata(path)
            kinds, _ = runs.setdefault(meta["run_id"] if meta else path, (set(), meta))
            kinds.add(kind)
    return runs


def cuda_release(meta):
    match = re.search(r"release ([\d.]+)", meta.get("installed_tools", {}).get("nvcc") or "")
    return f"CUDA {match.group(1)}" if match else "unknown"


def environment(manifest, rows):
    """GPU, host, driver, CUDA release, and source revision of the selected runs."""
    runs = selected_runs(rows)
    if not runs:
        return [manifest["label"], PENDING, PENDING, PENDING, PENDING]
    def values(get):
        return ", ".join(sorted({get(m) if m else "unrecorded" for _, m in runs.values()}))
    revision = manifest.get("source_revision")
    if not isinstance(revision, str) or not re.fullmatch(r"[0-9a-f]{40}", revision):
        raise ValueError(f"inputs/{manifest['folder']}.json: selected runs need a full "
                         f"source_revision commit hash")
    return [manifest["label"], values(lambda m: m["host"]), values(lambda m: m["gpu"]["driver"]),
            values(cuda_release), revision[:7]]


def short_kernel(name):
    """`<unnamed>::foo<128>(float *, int)` -> `foo<128>`."""
    for prefix in ("<unnamed>::", "(anonymous namespace)::"):
        name = name.replace(prefix, "")
    name = name.removeprefix("void ")
    base, bracket, rest = name.split("(")[0].partition("<")
    return base.split("::")[-1] + bracket + rest


def nsys_kernels(profile):
    """Per-invocation kernel time from an nsys kernel summary.

    Returns ([(kernel, us per invocation, launches per invocation)], [setup
    kernels]): setup kernels run a number of times that is not a multiple of
    the profiled invocations."""
    with open(REPO / profile["csv"]) as stream:
        rows = list(csv.DictReader(stream))
    kernels, setup = [], []
    for row in rows:
        count, total_us = int(row["Instances"]), float(row["Total Time (ns)"]) / 1e3
        if count % PROFILED_INVOCATIONS:
            setup.append(short_kernel(row["Name"]))
        else:
            kernels.append((short_kernel(row["Name"]), total_us / PROFILED_INVOCATIONS,
                            count // PROFILED_INVOCATIONS))
    return kernels, setup


def ncu_rows(profile):
    """All profiled launches of an ncu raw export, as {metric: (value, unit)} dicts."""
    with open(REPO / profile["csv"], newline="") as stream:
        rows = list(csv.reader(stream))
    header, units = rows[0], rows[1]
    return [{h: (v, u) for h, u, v in zip(header, units, row)} for row in rows[2:]]


def number(value):
    return float(value.replace(",", ""))


TIME_SCALE_TO_US = {"ns": 1e-3, "nsecond": 1e-3, "us": 1.0, "usecond": 1.0,
                    "ms": 1e3, "msecond": 1e3, "s": 1e6, "second": 1e6}


def launch_us(launch):
    value, unit = launch["gpu__time_duration.sum"]
    return number(value) * TIME_SCALE_TO_US[unit]


def ncu_metric_median(launches, metric):
    """Median within a configuration; missing or incompatible metrics stay blank."""
    samples = [launch.get(metric) for launch in launches]
    if any(sample is None or not sample[0] for sample in samples):
        return None
    if len({sample[1] for sample in samples}) != 1:
        return None
    try:
        return statistics.median(number(sample[0]) for sample in samples)
    except ValueError:
        return None


def ncu_dimensions(launch, kind):
    dims = [launch.get(f"launch__{kind}_dim_{axis}", ("", ""))[0] for axis in "xyz"]
    if all(dims):
        return " x ".join(dims)
    return launch.get(f"launch__{kind}_size", ("", ""))[0]


def repetitions(gpus):
    """The number of bench repetitions per selection, when all selections agree."""
    counts = {len(row[key]) for _, rows in gpus for row in rows for key in BENCHES if row[key]}
    return counts.pop() if len(counts) == 1 else None


def markdown(gpus):
    labels = [m["label"].split(" (")[0] for m, _ in gpus]
    runs = repetitions(gpus)
    lines = [
        f"# Private Loom CUDA workloads: results {VERSION}",
        "",
        "Timings and correctness checks for the six workloads in "
        "[`applications/loom_private`](../../applications/loom_private), ported from private "
        "Loom commit `4a75f8d`. Each workload has a baseline, the first direct port, and an "
        "optimized version of the same contract, which passes the same check. The older "
        "workloads have their own [results page](../README.md).",
        "",
        "- Times are microseconds per iteration, averaged over a back-to-back batch"
        + (f", then the median of {runs} recorded runs" if runs and runs > 1 else "")
        + ". Lower is better.",
        "- **Direct** launches every kernel on each iteration. **Graph** replays a CUDA Graph "
        "captured once. **Optimized** is the optimized version, replayed as a CUDA Graph.",
        "- *pending* means no run has been selected yet.",
        "",
        "## Workloads",
        "",
        "| Workload | Large problem |",
        "|---|---|",
    ]
    lines += [f"| {PROBLEMS[row['app']][0]} | {problem_text(row['app'], row['contract'])} |"
              for row in gpus[0][1]]
    scopes = [
        ("End-to-end time (µs)",
         "Includes copying inputs to the GPU and results back. The decoder keeps its weights on "
         "the GPU, and the anomaly model is loaded once, before timing. Each workload's "
         "`contract.json` defines the scope exactly.", "application"),
        ("Kernel-only time (µs)", "The same work without the copies.", "compute"),
    ]
    for title, note, scope in scopes:
        lines += ["", f"## {title}", "", note, "",
                  "| Workload | " + " | ".join(f"{g} Direct | {g} Graph | {g} Optimized"
                                             for g in labels) + " |",
                  "|---|" + "--:|--:|--:|" * len(labels)]
        for index, app in enumerate(APPS):
            cells = []
            for _, rows in gpus:
                cells += [fmt(timing(rows[index][key], scope)) for key in BENCHES]
            lines.append(f"| {PROBLEMS[app][0]} | " + " | ".join(cells) + " |")
    lines += ["", "## Optimized speedup", "",
              "Graph time divided by Optimized time, so both replay a CUDA Graph and only the "
              "implementation differs.", "",
              "| Workload | " + " | ".join(f"{g} end to end | {g} kernel only"
                                         for g in labels) + " |",
              "|---|" + "--:|--:|" * len(labels)]
    for index, app in enumerate(APPS):
        cells = []
        for _, rows in gpus:
            for scope in ("application", "compute"):
                graph = timing(rows[index]["bench_graph"], scope)
                optimized = timing(rows[index]["bench_optimized"], scope)
                cells.append(f"{graph / optimized:,.1f}×" if PENDING not in (graph, optimized)
                             else PENDING)
        lines.append(f"| {PROBLEMS[app][0]} | " + " | ".join(cells) + " |")
    lines += ["", "## Optimized versions", "",
              "Each workload's README explains its optimized version and the parameters chosen "
              "for each GPU.", ""]
    lines += [f"- **{PROBLEMS[app][0]}** ([details](../../applications/loom_private/{app}/"
              f"README.md#optimized-variant)): {OPTIMIZATIONS[app]}" for app in APPS]
    lines += ["", "## Correctness", "",
              "Each check compares every output with an independent CPU reference. The GPU runs "
              "twice with direct launches and twice with a graph, each time starting from "
              "poisoned buffers, and all four outputs must be identical. Checks are not "
              "timed. Each cell links to its record.", "",
              "| Workload | " + " | ".join(f"{g} baseline | {g} optimized" for g in labels)
              + " |",
              "|---|" + "---|---|" * len(labels)]
    for index, app in enumerate(APPS):
        lines.append(f"| {PROBLEMS[app][0]} | " + " | ".join(
            check_summary(rows[index][key]) for _, rows in gpus for key in CHECKS) + " |")
    lines += ["", "## Test machines", "",
              "Taken from the recorder metadata (`<gpu>/runs/<run>/metadata.json`) of the "
              "selected runs. Source revision is the code the runs were built from, as stated in "
              "`inputs/<gpu>.json`.", "",
              "| GPU | Host | Driver | CUDA | Source revision |", "|---|---|---|---|---|"]
    lines += ["| " + " | ".join(environment(m, rows)) + " |" for m, rows in gpus]
    lines += ["", "Generated by `results/loom_private/build_report.py` from the runs selected in "
              "`results/loom_private/inputs/<gpu>.json`. The workbook "
              f"[`loom_private_{VERSION}.xlsx`](loom_private_{VERSION}.xlsx) has the same data "
              "plus single-iteration latencies, per-kernel times and Nsight Compute metrics from "
              "the selected profiles, contract hashes, and per-invocation check results."]
    return "\n".join(lines)


# --------------------------------------------------------------- workbook ---
# The workbook uses the layout and styling of results/cuda_performance.xlsx
# (results/build_spreadsheet.py). The helpers are copied rather than imported so
# that this report never loads or validates the legacy inputs.

GITHUB = "https://github.com/ankaijin/cuda-elastic/blob"
BASELINE = "rtx3090"
VERSIONS = ["Direct", "Graph", "Optimized"]
STAGES = [("application", "Application (end to end)"), ("compute", "Kernel only (no transfers)")]
METRIC = "CUDA events (batched)"
BENCH_PENDING = "pending (no bench run selected)"
NSYS_PENDING = "pending (no Nsight Systems profile selected)"
NCU_PENDING = "pending (no Nsight Compute profile selected)"
# Short visible note per application and stage; the contract's full timing
# scope goes in the cell comment.
SHORT_NOTES = {
    "application": {
        "decoder": "Weights stay on the GPU; copies prompts in and tokens out.",
        "anomaly": "Model loaded before timing; copies samples in and outputs out.",
    },
    "compute": {},
}
DEFAULT_NOTES = {"application": "Includes input upload and output download (pinned host memory).",
                 "compute": "Kernels only; inputs already on the GPU."}
# Per-launch ncu metrics on the NCU launches sheets (metric, column label).
NCU_DETAIL_METRICS = [
    ("gpu__time_duration.sum", "Duration"),
    ("sm__throughput.avg.pct_of_peak_sustained_elapsed", "SM throughput %"),
    ("gpu__compute_memory_throughput.avg.pct_of_peak_sustained_elapsed", "Memory throughput %"),
    ("gpu__dram_throughput.avg.pct_of_peak_sustained_elapsed", "DRAM throughput %"),
    ("sm__warps_active.avg.pct_of_peak_sustained_active", "Achieved occupancy %"),
    ("sm__maximum_warps_per_active_cycle_pct", "Theoretical occupancy %"),
    ("launch__registers_per_thread", "Registers / thread"),
    ("launch__grid_size", "Grid"),
    ("launch__block_size", "Block"),
    ("launch__waves_per_multiprocessor", "Waves / SM"),
]
# Launch configuration: kernels with different values are separate rows.
NCU_CONFIG_METRICS = ["launch__grid_size", "launch__block_size", "launch__registers_per_thread",
                      "launch__shared_mem_per_block"]


def link(path):
    """GitHub URL for path, pinned to the latest commit on origin/main that has
    the file; files not merged yet link to main."""
    log = subprocess.run(["git", "-C", str(REPO), "log", "--format=%H", "-1", "origin/main",
                          "--", path], capture_output=True, text=True).stdout.strip()
    return f"{GITHUB}/{log or 'main'}/{path}"


def workbook(gpus):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    blue = Font(color="0563C1", underline="single")
    bold = Font(bold=True)
    grey = Font(italic=True, color="808080")
    header_fill = PatternFill("solid", fgColor="1F3864")
    header_font = Font(bold=True, color="FFFFFF")
    fills = [PatternFill("solid", fgColor=c) for c in ("FCE4D6", "E2EFDA", "DDEBF7")]
    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    links = {}
    book = Workbook()
    book.remove(book.active)

    def github(path):
        if path not in links:
            links[path] = link(path)
        return links[path]

    def table_sheet(title, headers, widths, description):
        ws = book.create_sheet(title)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        ws.cell(1, 1, title).font = Font(bold=True, size=14)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
        ws.cell(2, 1, description).alignment = wrap
        ws.row_dimensions[2].height = 42
        for i, (header, width) in enumerate(zip(headers, widths), 1):
            c = ws.cell(4, i, header)
            c.font, c.fill, c.border = header_font, header_fill, box
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.row_dimensions[4].height = 42
        ws.freeze_panes = "C5"
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.print_title_rows = "1:4"
        return ws

    def finish_table(ws, height=58):
        ws.auto_filter.ref = f"A4:{get_column_letter(ws.max_column)}{ws.max_row}"
        for row in ws.iter_rows(min_row=5):
            for c in row:
                c.alignment = wrap
                c.border = box
            ws.row_dimensions[row[0].row].height = height

    def nav(ws, targets):
        for col, target in enumerate(targets, 1):
            c = ws.cell(3, col, target)
            c.hyperlink, c.font = f"#'{target}'!A1", blue

    def value(ws, row, col, data, source=None, speedup=False):
        """A number (linked to its source) or grey pending text with its reason."""
        c = ws.cell(row, col)
        if isinstance(data, (int, float)):
            c.value = round(data, 3)
            c.number_format = '0.00"x"' if speedup else '#,##0.00'
            if source:
                c.hyperlink, c.font = github(source), blue
        elif data is not None:
            c.value = "pending" if data.startswith("pending") else data
            c.font = grey
            if data.startswith("pending"):
                c.comment = Comment(data, "build_report.py")
        return c

    def put_link(ws, row, col, path, text):
        c = ws.cell(row, col, text)
        c.hyperlink, c.font = github(path), blue
        return c

    def note(ws, row, col, short, full):
        c = ws.cell(row, col, short)
        c.comment = Comment(full, "build_report.py", width=480, height=280)
        return c

    def ratio(a, b):
        numbers = isinstance(a, (int, float)) and isinstance(b, (int, float))
        return a / b if numbers and b else None

    def bench(records, scope, field="batch_per_iteration_us"):
        return BENCH_PENDING if not records else timing(records, scope, field)

    def bench_source(records, scope):
        return median_record(records, scope)["_path"] if records else None

    def kernel_time(profile):
        return NSYS_PENDING if profile is None else sum(us for _, us, _ in nsys_kernels(profile)[0])

    def describe(records):
        counts = {r["iterations"] for r in records}
        runs = len(records)
        text = f" {counts.pop():,} iterations/batch." if len(counts) == 1 else ""
        return text + (" One run." if runs == 1 else f" Median of {runs} runs.")

    def stage_note(row, scope):
        spec = row["contract"] or {}
        short = SHORT_NOTES[scope].get(row["app"], DEFAULT_NOTES[scope])
        records = [r for key in BENCHES for r in row[key] or []]
        if records:
            short += describe(records)
        full = spec.get("timing_scope", {}).get(scope, "Contract pending.")
        return short, full[0].upper() + full[1:] + "."

    def kernel_note(row):
        parts, setup_kernels = [], set()
        for version, (kind, key) in zip(VERSIONS, [("nsys", v) for v in PROFILE_VERSIONS]):
            profile = row["profiles"][(kind, key)]
            if profile is None:
                continue
            kernels, setup = nsys_kernels(profile)
            setup_kernels.update(setup)
            total = sum(us for _, us, _ in kernels) or 1
            parts.append(f"{version}: " + ", ".join(
                f"{name} ×{count} ({100 * us / total:.0f}%)"
                for name, us, count in sorted(kernels, key=lambda k: -k[1])))
        if not parts:
            return "Not profiled yet.", "No Nsight Systems profile is selected."
        launches = sum(c for _, _, c in nsys_kernels(row["profiles"][("nsys", "optimized")]
                                                     or row["profiles"][("nsys", "graph")]
                                                     or row["profiles"][("nsys", "direct")])[0])
        short = f"Sum of every kernel in one invocation ({launches:,} launches, last version)."
        full = ("Kernel time per invocation from Nsight Systems: the profile runs four "
                "invocations (two end to end, two kernel only), and each kernel's total is "
                "divided by four. " + " ".join(f"{p}." for p in parts))
        if setup_kernels:
            full += " Setup kernels, not counted: " + ", ".join(sorted(setup_kernels)) + "."
        return short, full

    names = {m["folder"]: m["label"].split(" (")[0] for m, _ in gpus}
    base_rows = dict((m["folder"], rows) for m, rows in gpus)[BASELINE]
    targets = [m["folder"] for m, _ in gpus if m["folder"] != BASELINE]
    shared = ["Sources", "GPU analysis", "Correctness", "Environment", "Contracts"]

    def summary_headers(folder, first):
        title = names[folder]
        headers = ["Workload", "Stage", "Problem size", "Metric", "Direct (us)", "Graph (us)",
                   "Optimized (us)", "Graph speedup\nDirect / Graph",
                   "Optimized speedup\nGraph / Optimized"]
        if folder != BASELINE:
            headers += [f"GPU speedup\n{names[BASELINE].removeprefix('RTX ')} / "
                        f"{title.removeprefix('RTX ')}\n(Graph)"]
        widths = [24, first, 30, 20, 15, 15, 15, 15, 17] + ([18] if folder != BASELINE else [])
        return headers + ["Notes"], widths + [48]

    def speedup(a, b):
        """a / b, else the pending reason of the missing value."""
        return ratio(a, b) or next((x for x in (b, a) if isinstance(x, str)), BENCH_PENDING)

    def summary_values(ws, r, direct, graph, optimized, sources, base_graph=None, compare=False):
        for col, (data, source, fill) in enumerate(zip((direct, graph, optimized), sources,
                                                       fills), 5):
            value(ws, r, col, data, source).fill = fill
        value(ws, r, 8, speedup(direct, graph), speedup=True)
        value(ws, r, 9, speedup(graph, optimized), speedup=True)
        if compare:
            value(ws, r, 10, speedup(base_graph, graph), speedup=True)

    for manifest, rows in gpus:
        folder, title = manifest["folder"], names[manifest["folder"]]
        headers, widths = summary_headers(folder, 30)
        ws = table_sheet(title, headers, widths,
                         f"{manifest['label']}; {manifest['toolkit']}. Large workloads: time per "
                         "iteration in a back-to-back batch; lower is better. Optimized is the "
                         "optimized version of the same contract, replayed as a CUDA Graph. Times "
                         "link to the bench records. Notes have full explanations in comments.")
        nav(ws, [f"Latency {title}", f"Kernels {title}"] + shared)
        for index, row in enumerate(rows):
            for scope, stage in STAGES:
                r = ws.max_row + 1
                for col, text in enumerate((PROBLEMS[row["app"]][0], stage,
                                            problem_text(row["app"], row["contract"]), METRIC), 1):
                    ws.cell(r, col, text)
                ws.cell(r, 1).font = bold
                summary_values(ws, r, *(bench(row[key], scope) for key in BENCHES),
                               [bench_source(row[key], scope) for key in BENCHES],
                               bench(base_rows[index]["bench_graph"], scope), folder != BASELINE)
                note(ws, r, len(headers), *stage_note(row, scope))
        finish_table(ws)

    for manifest, rows in gpus:
        folder, title = manifest["folder"], names[manifest["folder"]]
        headers = ["Workload", "Stage", "Problem size", "Version", "Batch iterations", "Runs",
                   "Full batch (ms)", "Batched (us/iteration)", "Single (us)", "Single / batched",
                   "Graph speedup (single)", "Optimized speedup (single)", "Notes"]
        ws = table_sheet(f"Latency {title}", headers,
                         [24, 24, 30, 16, 14, 9, 16, 17, 15, 15, 17, 18, 48],
                         f"{manifest['label']}. Single = median of the isolated iterations that "
                         "follow the batch (as many as the batch has). Single/batched reflects "
                         "scheduling and timing conditions, not pure launch overhead. With "
                         "several runs, each column is the median over the runs; missing "
                         "measurements stay pending.")
        nav(ws, [title, f"Kernels {title}"] + shared)
        for row in rows:
            for scope, stage in STAGES:
                singles = {}
                for version, key, fill in zip(VERSIONS, BENCHES, fills):
                    records = row[key]
                    r = ws.max_row + 1
                    for col, text in enumerate((PROBLEMS[row["app"]][0], stage,
                                                problem_text(row["app"], row["contract"]),
                                                version), 1):
                        ws.cell(r, col, text)
                    ws.cell(r, 1).font = bold
                    if not records:
                        for col in range(5, 11):
                            value(ws, r, col, BENCH_PENDING, speedup=col == 10)
                    else:
                        source = bench_source(records, scope)
                        singles[version] = timing(records, scope, "single_median_us")
                        iterations = {rec["iterations"] for rec in records}
                        ws.cell(r, 5, iterations.pop() if len(iterations) == 1 else "mixed")
                        ws.cell(r, 6, len(records))
                        for col, field in enumerate(("batch_total_ms", "batch_per_iteration_us",
                                                     "single_median_us"), 7):
                            value(ws, r, col, timing(records, scope, field), source)
                        value(ws, r, 10, ratio(singles[version],
                                               timing(records, scope)), speedup=True)
                    if version == "Graph":
                        value(ws, r, 11, ratio(singles.get("Direct"), singles.get("Graph"))
                              or BENCH_PENDING, speedup=True)
                    if version == "Optimized":
                        value(ws, r, 12, ratio(singles.get("Graph"), singles.get("Optimized"))
                              or BENCH_PENDING, speedup=True)
                    note(ws, r, 13, *stage_note(row, scope))
                    for col in range(1, len(headers) + 1):
                        ws.cell(r, col).fill = fill
        finish_table(ws)

        headers, widths = summary_headers(folder, 26)
        ws = table_sheet(f"Kernels {title}", headers, widths,
                         f"{manifest['label']}; {manifest['toolkit']}. GPU kernel time per "
                         "invocation from Nsight Systems: every kernel of one invocation summed, "
                         "with the copies left out. Times link to the profiles; notes list the "
                         "kernels and their shares.")
        nav(ws, [title, f"Latency {title}"] + shared)
        for index, row in enumerate(rows):
            r = ws.max_row + 1
            for col, text in enumerate((PROBLEMS[row["app"]][0], "All kernels, one invocation",
                                        problem_text(row["app"], row["contract"]),
                                        "Nsight Systems kernel sum"), 1):
                ws.cell(r, col, text)
            ws.cell(r, 1).font = bold
            profiles = [row["profiles"][("nsys", v)] for v in PROFILE_VERSIONS]
            summary_values(ws, r, *(kernel_time(p) for p in profiles),
                           [p["_path"] if p else None for p in profiles],
                           kernel_time(base_rows[index]["profiles"][("nsys", "graph")]),
                           folder != BASELINE)
            note(ws, r, len(headers), *kernel_note(row))
        finish_table(ws)

    ws = table_sheet("Sources", ["GPU", "Workload", "Problem size", "Version", "Measurement",
                                 "Run metadata", "NCU report", "Nsight Systems", "Notes"],
                     [15, 24, 30, 15, 44, 44, 44, 44, 48],
                     "Source files for every timing. Filenames link to the original evidence: "
                     "the bench record with the median end-to-end time, its run's metadata, and "
                     "the selected profiles.")
    nav(ws, [names[f] for f in GPUS if f in names] + shared[1:])
    for manifest, rows in gpus:
        for row in rows:
            for version, key, profile in zip(VERSIONS, BENCHES, PROFILE_VERSIONS):
                records = row[key]
                r = ws.max_row + 1
                for col, text in enumerate((names[manifest["folder"]], PROBLEMS[row["app"]][0],
                                            problem_text(row["app"], row["contract"]), version), 1):
                    ws.cell(r, col, text)
                if not records:
                    value(ws, r, 5, BENCH_PENDING)
                else:
                    record = median_record(records, "application")
                    runs = f" (median of {len(records)})" if len(records) > 1 else ""
                    put_link(ws, r, 5, record["_path"], Path(record["_path"]).name + runs)
                    meta = run_metadata(record["_path"])
                    if meta:
                        path = f"{manifest['folder']}/runs/{meta['run_id']}/metadata.json"
                        put_link(ws, r, 6, path, f"{meta['run_id']}/metadata.json")
                for col, kind, pending in ((7, "ncu", NCU_PENDING), (8, "nsys", NSYS_PENDING)):
                    selected = row["profiles"][(kind, profile)]
                    if selected:
                        put_link(ws, r, col, selected["_path"], Path(selected["_path"]).name)
                    else:
                        value(ws, r, col, manifest.get(f"{kind}_pending", pending))
                note(ws, r, 9, *stage_note(row, "application"))
                for col in range(1, 10):
                    ws.cell(r, col).fill = fills[VERSIONS.index(version)]
    finish_table(ws)

    ws = table_sheet("GPU analysis", ["GPU", "Workload", "Stage", "Problem size", "Version",
                                      f"{names[BASELINE]} (us)", "GPU (us)", "Measured speedup",
                                      "Notes"],
                     [15, 24, 26, 30, 14, 15, 15, 16, 48],
                     f"Large workloads only. Measured speedup = {names[BASELINE]} time / target "
                     "GPU time for the same version, scope and source revision. No predictions: "
                     "the hardware models of the legacy report do not describe these workloads.")
    nav(ws, [names[f] for f in GPUS if f in names] + ["Sources"])
    for manifest, rows in gpus:
        if manifest["folder"] == BASELINE:
            continue
        for index, row in enumerate(rows):
            for scope, stage in STAGES:
                for version, key, fill in zip(VERSIONS[1:], BENCHES[1:], fills[1:]):
                    r = ws.max_row + 1
                    for col, text in enumerate((names[manifest["folder"]],
                                                PROBLEMS[row["app"]][0], stage,
                                                problem_text(row["app"], row["contract"]),
                                                version), 1):
                        ws.cell(r, col, text)
                    ws.cell(r, 2).font = bold
                    base, target = bench(base_rows[index][key], scope), bench(row[key], scope)
                    value(ws, r, 6, base, bench_source(base_rows[index][key], scope))
                    value(ws, r, 7, target, bench_source(row[key], scope))
                    value(ws, r, 8, ratio(base, target) or BENCH_PENDING, speedup=True)
                    note(ws, r, 9, *stage_note(row, scope))
                    for col in range(1, 10):
                        ws.cell(r, col).fill = fill
    if not targets:
        value(ws, 5, 1, "pending (only one GPU)")
    finish_table(ws)

    for manifest, rows in gpus:
        title = names[manifest["folder"]]
        metrics, raw_title = f"ncu metrics ({title})", f"NCU launches {title}"
        headers = ["Workload", "Version", "Kernel", "Launches", "Median (us)", "Min (us)",
                   "Max (us)", "SM throughput %\n(median)", "DRAM throughput %\n(median)",
                   "Achieved occupancy %\n(median)", "Memory throughput %\n(median)",
                   "Theoretical occupancy %\n(median)", "Grid (x y z)", "Block (x y z)",
                   "Registers / thread", "Shared memory / block", "ncu report", "Launch details",
                   "Problem size"]
        ws = table_sheet(metrics, headers,
                         [24, 14, 46, 11, 15, 15, 15, 17, 17, 18, 18, 19, 22, 22, 17, 22, 50, 20,
                          30],
                         "Large workloads; one row per report, kernel, and launch configuration. "
                         "Each report covers the setup kernels and one end-to-end invocation. "
                         "Medians and min/max describe profiled launches, not application "
                         "latency or total GPU utilization. Missing metrics stay blank.")
        raw = table_sheet(raw_title, ["Workload", "Version", "ncu report", "Launch #", "Kernel"]
                          + [label for _, label in NCU_DETAIL_METRICS],
                          [24, 14, 50, 11, 46] + [18] * len(NCU_DETAIL_METRICS),
                          "Every profiled launch in the selected reports, grouped by kernel and "
                          "configuration; Launch # is the original report order (from zero). "
                          "Durations and units are as Nsight Compute reports them.")
        ws.freeze_panes = raw.freeze_panes = "D5"
        nav(ws, [raw_title, title])
        nav(raw, [metrics, title])
        reports = [(row, version, row["profiles"][("ncu", key)])
                   for row in rows for version, key in zip(VERSIONS, PROFILE_VERSIONS)
                   if row["profiles"][("ncu", key)]]
        if not reports:
            for sheet in (ws, raw):
                value(sheet, 5, 1, manifest.get("ncu_pending", NCU_PENDING))
        for row, version, profile in reports:
            fill = fills[VERSIONS.index(version)]
            workload = PROBLEMS[row["app"]][0]
            groups = {}
            for i, launch in enumerate(ncu_rows(profile)):
                key = (launch["Kernel Name"][0],
                       tuple(launch.get(metric) for metric in NCU_CONFIG_METRICS))
                groups.setdefault(key, []).append((i, launch))
            for (kernel, _), indexed in groups.items():
                launches = [launch for _, launch in indexed]
                first = launches[0]
                durations = [launch_us(launch) for launch in launches]
                r, first_raw = ws.max_row + 1, raw.max_row + 1
                values = [workload, version, short_kernel(kernel), len(launches),
                          statistics.median(durations), min(durations), max(durations)]
                values += [ncu_metric_median(launches, metric) for metric in (
                    "sm__throughput.avg.pct_of_peak_sustained_elapsed",
                    "gpu__dram_throughput.avg.pct_of_peak_sustained_elapsed",
                    "sm__warps_active.avg.pct_of_peak_sustained_active",
                    "gpu__compute_memory_throughput.avg.pct_of_peak_sustained_elapsed",
                    "sm__maximum_warps_per_active_cycle_pct")]
                values += [ncu_dimensions(first, "grid"), ncu_dimensions(first, "block"),
                           ncu_metric_median(launches, "launch__registers_per_thread")]
                shared_mem, unit = first.get("launch__shared_mem_per_block", ("", ""))
                values += [f"{shared_mem} {unit}".strip()]
                for col, data in enumerate(values, 1):
                    if isinstance(data, (int, float)) and col > 4:
                        value(ws, r, col, data)
                    else:
                        ws.cell(r, col, data)
                ws.cell(r, 3).comment = Comment(kernel, "build_report.py", width=540, height=160)
                put_link(ws, r, 17, profile["_path"], Path(profile["_path"]).name)
                c = ws.cell(r, 18, f"{len(launches)} launches")
                c.hyperlink, c.font = f"#'{raw_title}'!A{first_raw}", blue
                ws.cell(r, 19, problem_text(row["app"], row["contract"]))
                for col in range(1, len(headers) + 1):
                    ws.cell(r, col).fill = fill
                for i, launch in indexed:
                    rr = raw.max_row + 1
                    raw.cell(rr, 1, workload)
                    raw.cell(rr, 2, version)
                    put_link(raw, rr, 3, profile["_path"], Path(profile["_path"]).name)
                    raw.cell(rr, 4, i)
                    raw.cell(rr, 5, short_kernel(launch["Kernel Name"][0]))
                    for col, (metric, _) in enumerate(NCU_DETAIL_METRICS, 6):
                        if metric not in launch:
                            continue
                        data, unit = launch[metric]
                        try:
                            data = number(data)
                        except ValueError:
                            pass
                        raw.cell(rr, col, f"{data} {unit}".strip() if unit and unit != "%"
                                 else data)
                    for col in range(1, raw.max_column + 1):
                        raw.cell(rr, col).fill = fill
        finish_table(ws)
        finish_table(raw, height=30)

    ws = table_sheet("Correctness", ["GPU", "Workload", "Version", "Contract SHA-256", "Status",
                                     "Direct", "Graph", "Direct = Graph", "Record"],
                     [15, 24, 14, 30, 12, 12, 12, 14, 44],
                     "Large-workload checks, not timings. Each runs the CPU reference once, then "
                     "two Direct and two Graph invocations from poisoned buffers; every output "
                     "must match the reference and all four must be identical. The baseline and "
                     "the optimized version are checked separately against the same contract.")
    for manifest, rows in gpus:
        for row in rows:
            for version, key in (("Baseline", "check"), ("Optimized", "check_optimized")):
                r = ws.max_row + 1
                spec, record = row["contract"] or {}, row[key]
                for col, text in enumerate((names[manifest["folder"]], PROBLEMS[row["app"]][0],
                                            version, spec.get("_sha256", PENDING)), 1):
                    ws.cell(r, col, text)
                for col, text in enumerate(check_cells(record), 5):
                    if text == PENDING:
                        value(ws, r, col, "pending (no check run selected)")
                    else:
                        ws.cell(r, col, text)
                if record:
                    put_link(ws, r, 9, record["_path"], Path(record["_path"]).name)
    finish_table(ws)

    ws = table_sheet("Environment", ["GPU", "Run", "Measurement", "Host", "Driver", "CUDA",
                                     "Source revision"],
                     [15, 44, 20, 12, 14, 12, 16],
                     "Recorder metadata of every selected run. Source revision is the code the "
                     "runs were built from, as stated in results/loom_private/inputs/<gpu>.json.")
    for manifest, rows in gpus:
        for run_id, (kinds, meta) in sorted(selected_runs(rows).items()):
            if meta is None:
                continue
            r = ws.max_row + 1
            ws.cell(r, 1, names[manifest["folder"]])
            put_link(ws, r, 2, f"{manifest['folder']}/runs/{run_id}/metadata.json", run_id)
            for col, text in enumerate((", ".join(sorted(kinds)), meta["host"],
                                        meta["gpu"]["driver"], cuda_release(meta),
                                        manifest["source_revision"][:7]), 3):
                ws.cell(r, col, text)
    finish_table(ws)

    ws = table_sheet("Contracts", ["Workload", "Contract", "Contract SHA-256", "Upstream commit",
                                   "Extension of upstream"],
                     [24, 44, 30, 30, 60],
                     "Each workload's contract.json fixes its dimensions, inputs, outputs, "
                     "tolerances, and timing scopes. Every result record names the contract hash "
                     "it validated; the baseline and the optimized version share it.")
    for app in APPS:
        spec = contract(app) or {}
        r = ws.max_row + 1
        ws.cell(r, 1, PROBLEMS[app][0]).font = bold
        put_link(ws, r, 2, f"applications/loom_private/{app}/contract.json",
                 f"{app}/contract.json")
        ws.cell(r, 3, spec.get("_sha256", PENDING))
        ws.cell(r, 4, spec.get("upstream", {}).get("commit", PENDING))
        extension = spec.get("extension_of_upstream", PENDING)
        ws.cell(r, 5, "No (direct port)" if extension is False else str(extension))
    finish_table(ws)

    # Legacy order: summaries, then Latency and Kernels per GPU, Sources, GPU
    # analysis, and the ncu sheets; this report's own sheets come last.
    order = [names[f] for f in GPUS if f in names]
    order += [f"{kind} {names[f]}" for f in GPUS if f in names for kind in ("Latency", "Kernels")]
    order += ["Sources", "GPU analysis"]
    order += [s for f in GPUS if f in names
              for s in (f"ncu metrics ({names[f]})", f"NCU launches {names[f]}")]
    order += ["Correctness", "Environment", "Contracts"]
    book._sheets = [book[title] for title in order]
    return book


def main():
    gpus = [load_gpu(folder) for folder in GPUS]
    MD_OUT.write_text(markdown(gpus) + "\n")
    workbook(gpus).save(XLSX_OUT)
    print(f"wrote {MD_OUT.relative_to(REPO)} and {XLSX_OUT.relative_to(REPO)}")


if __name__ == "__main__":
    main()
