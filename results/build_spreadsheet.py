#!/usr/bin/env python3
"""Build results/cuda_performance.xlsx and results/README.md from existing profiles.

The RTX 3090, RTX 4090, and RTX 5090 summary sheets compare application microseconds per
iteration in a back-to-back batch: Direct, Graph, and Agentic optimized, plus
the two implementation speedups. Target GPUs also show measured GPU speedup.
Each timing links to its original measurement. Missing results stay pending.

Supporting sheets retain full-batch and isolated timings, batch/run counts,
single/batched ratios, isolated-latency speedups, kernel comparisons, source
files, GPU predictions, and every profiled NCU launch. Visible notes are short;
comments retain full explanations and prediction assumptions.

Inputs are pinned explicitly in results/inputs/<gpu>.json. Null means pending;
selected missing files are errors. All data comes from each GPU's existing
output folder; no benchmark is rerun. RTX 5090 predictions remain disabled.
Whole-application timings use time_all.sh logs when available; otherwise they
use benchmark CUDA events or the median iteration span in a Nsight Systems
trace. Kernel rows use NCU durations, with optimized speedups compared against
the baseline kernels replaced. Graph = Direct / Graph; optimized = Graph /
Optimized. Single/batched ratios include scheduling and measurement effects
and do not isolate launch overhead.

Requires ncu (Nsight Compute CLI) on PATH and openpyxl.
Usage: python3 results/build_spreadsheet.py
"""

import bisect
import csv
from functools import lru_cache
import io
import json
import re
import sqlite3
import statistics
import subprocess
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parent.parent
OUT = REPO / "results" / "cuda_performance.xlsx"
MD_OUT = REPO / "results" / "README.md"
GITHUB = "https://github.com/ankaijin/cuda-elastic/blob"


# ------------------------------------------------------------------- GPUs ---

class GPU:
    """One GPU's results and exact, GPU-relative input paths; null is pending."""

    def __init__(self, folder, label, toolkit, specs, files, notes):
        if folder not in ("rtx3090", "rtx4090", "rtx5090"):
            raise ValueError(f"Unsupported GPU folder: {folder}")
        self.folder, self.label, self.toolkit = folder, label, toolkit
        self.specs, self.files, self.notes = specs, files, notes

    @property
    def sheet(self):
        return self.label.split(" (")[0]

    def file(self, key):
        selected = self.files.get(key)
        if selected is None:
            return None
        if (not isinstance(selected, str) or not selected
                or any(c in selected for c in "*?[")
                or Path(selected).is_absolute() or ".." in Path(selected).parts):
            raise ValueError(f"{self.folder}: {key} must be an exact GPU-relative path")
        root = (REPO / self.folder).resolve()
        path = REPO / self.folder / selected
        if not path.resolve().is_relative_to(root):
            raise ValueError(f"{self.folder}: {key} escapes the GPU directory")
        if not path.is_file():
            raise FileNotFoundError(f"{self.folder}: selected {key} is missing: {path}")
        return str(path.relative_to(REPO))


MODES = ("graph", "direct")
# Existing optimized implementations' launch modes.
SOTA_LAUNCH = {"llama": "graph", "loom": "direct", "pagerank": "direct",
               "anomaly": "graph", "kws": "direct", "vww": "direct"}


# time_all.sh's small shapes (SHAPES=small), under input keys <key>_small.
SMALL = "_small"
SMALL_TIMING_NAMES = {
    "llama_small": ("llama2c", "llama2c_small", "llama2c_sota_small"),
    "loom_small": ("loom_msa", "msa_optimized_4x4", "msa_sota_4x4"),
    "pagerank_small": ("pagerank", "pagerank_small", "pagerank_sota_small"),
    "anomaly_small": ("mlperf-anomaly", "anomaly_mse_small", "anomaly_mse_sota_small"),
    "kws_small": ("mlperf-classification", "kws_postprocess_small", "kws_postprocess_sota_small"),
    "vww_small": ("mlperf-classification", "vww_preprocess_small", "vww_preprocess_sota_small"),
}
SINGLE_ITERATIONS = 20  # each harness's kSingleIterations / kSingleReplays



def load_gpus():
    """Load each GPU's independently maintained, explicit input manifest."""
    gpus = []
    for folder in ("rtx3090", "rtx4090", "rtx5090"):
        data = json.loads((REPO / "results" / "inputs" / f"{folder}.json").read_text())
        if data["folder"] != folder:
            raise ValueError(f"{folder}.json must select the {folder} output directory")
        gpus.append(GPU(**data))
    return gpus


def validate_inputs(gpus):
    """Validate every selection before any report is generated, including fallbacks."""
    for gpu in gpus:
        for key in gpu.files:
            gpu.file(key)


GPUS = load_gpus()
RTX3090, RTX4090, RTX5090 = GPUS
BASELINE = RTX3090


# ---------------------------------------------------------------- loaders ---

def git(*args):
    return subprocess.run(["git", "-C", str(REPO), *args], capture_output=True, text=True)


@lru_cache(maxsize=None)
def link(path):
    """GitHub URL for path, pinned to the latest commit on origin/main that
    has the file, under the name it had there (files moved since, e.g. into
    rtx3090/, keep their original links); files not merged yet link to main."""
    log = git("log", "--follow", "--name-only", "--format=%x00%H", "--", path).stdout
    for entry in log.split("\0")[1:]:
        commit, _, name = entry.strip().partition("\n")
        if git("merge-base", "--is-ancestor", commit, "origin/main").returncode == 0:
            return f"{GITHUB}/{commit}/{name.strip()}"
    return f"{GITHUB}/main/{path}"


_ncu_cache = {}


def ncu_rows(path):
    """All profiled launches in an ncu report, as {metric: (value, unit)} dicts."""
    if path not in _ncu_cache:
        text = subprocess.run(["ncu", "--import", str(REPO / path), "--csv", "--page", "raw"],
                              check=True, capture_output=True, text=True).stdout
        rows = list(csv.reader(io.StringIO(text)))
        header, units = rows[0], rows[1]
        _ncu_cache[path] = [
            {h: (v, u) for h, u, v in zip(header, units, row)} for row in rows[2:]
        ]
    return _ncu_cache[path]


def number(value):
    return float(value.replace(",", ""))


TIME_SCALE_TO_US = {"ns": 1e-3, "nsecond": 1e-3, "us": 1.0, "usecond": 1.0,
                    "ms": 1e3, "msecond": 1e3, "s": 1e6, "second": 1e6}


def kernel_name(row):
    return row["Kernel Name"][0].split("(")[0].split("::")[-1].removeprefix("void ")


def launch_us(row):
    v, unit = row["gpu__time_duration.sum"]
    return number(v) * TIME_SCALE_TO_US[unit]


def kernel_launches(path, kernel, grid=None):
    """Launches of `kernel` (name prefix after any namespace) in an ncu report."""
    return [r for r in ncu_rows(path)
            if kernel_name(r).startswith(kernel)
            and (grid is None or number(r["launch__grid_size"][0]) == grid)]


def ncu_kernel_us(path, kernel, grid=None):
    launches = kernel_launches(path, kernel, grid)
    if not launches:
        raise ValueError(f"{kernel} (grid={grid}) not found in {path}")
    times = [launch_us(r) for r in launches]
    return statistics.mean(times), len(times)


def ncu_time_shares(path):
    """Percent of the report's summed kernel time per kernel name."""
    totals = {}
    for r in ncu_rows(path):
        totals[kernel_name(r)] = totals.get(kernel_name(r), 0) + launch_us(r)
    return {k: 100 * v / sum(totals.values()) for k, v in totals.items()}


def nsys_node_replay_us(path):
    """Median span of CUDA Graph replays that include memcpy nodes, in an nsys
    trace taken with --cuda-graph-trace=node (no per-graph trace rows)."""
    with sqlite3.connect(REPO / path) as db:
        launches = {r[0] for r in db.execute(
            "SELECT r.correlationId FROM CUPTI_ACTIVITY_KIND_RUNTIME r "
            "JOIN StringIds s ON s.id = r.nameId WHERE s.value LIKE 'cudaGraphLaunch%'")}
        spans = {}
        for table in ("CUPTI_ACTIVITY_KIND_KERNEL", "CUPTI_ACTIVITY_KIND_MEMCPY"):
            for corr, start, end in db.execute(
                    f"SELECT correlationId, start, end FROM {table} WHERE graphNodeId IS NOT NULL"):
                span = spans.setdefault(corr, [start, end, False])
                span[0], span[1] = min(span[0], start), max(span[1], end)
                span[2] |= table.endswith("MEMCPY")
    durations = [(end - start) / 1e3 for corr, (start, end, copies) in spans.items()
                 if corr in launches and copies]
    return statistics.median(durations), len(durations)


def nsys_nvtx_span_us(path, range_name):
    """Median GPU span (first start to last end) of the work enqueued inside
    each NVTX push/pop range named range_name, in an nsys trace taken with
    --trace=cuda,nvtx, and the number of ranges; None when the trace has no
    such range. A range's work is every kernel, memcpy and whole-graph
    execution whose correlationId belongs to a runtime API call made on the
    range's thread while it was open: the launches and copies of a direct
    iteration, or a cudaGraphLaunch (its nodes with --cuda-graph-trace=node,
    the whole graph with =graph)."""
    with sqlite3.connect(REPO / path) as db:
        tables = {r[0] for r in db.execute("SELECT name FROM sqlite_master WHERE type = 'table'")}
        if "NVTX_EVENTS" not in tables:
            return None
        ranges = db.execute(
            "SELECT n.start, n.end, n.globalTid FROM NVTX_EVENTS n "
            "LEFT JOIN StringIds s ON s.id = n.textId "
            "WHERE n.eventType = 59 AND COALESCE(n.text, s.value) = ?", (range_name,)).fetchall()
        if not ranges:
            return None
        calls = {}
        for start, tid, corr in db.execute(
                "SELECT start, globalTid, correlationId FROM CUPTI_ACTIVITY_KIND_RUNTIME "
                "ORDER BY start"):
            starts, corrs = calls.setdefault(tid, ([], []))
            starts.append(start)
            corrs.append(corr)
        work = {}
        for table in ("CUPTI_ACTIVITY_KIND_KERNEL", "CUPTI_ACTIVITY_KIND_MEMCPY",
                      "CUPTI_ACTIVITY_KIND_GRAPH_TRACE"):
            if table in tables:
                for corr, start, end in db.execute(
                        f"SELECT correlationId, start, end FROM {table}"):
                    work.setdefault(corr, []).append((start, end))
    spans = []
    for start, end, tid in ranges:
        starts, corrs = calls.get(tid, ([], []))
        first, last = bisect.bisect_left(starts, start), bisect.bisect_right(starts, end)
        intervals = [w for corr in corrs[first:last] for w in work.get(corr, [])]
        if intervals:
            spans.append((max(e for _, e in intervals) - min(s for s, _ in intervals)) / 1e3)
    return statistics.median(spans), len(spans)


def nsys_iteration(path, range_name):
    """(median us, count, unit) of one end-to-end iteration in an nsys trace:
    the NVTX ranges named range_name where the trace has them, else (traces
    from before the harnesses had NVTX ranges) the graph replays with copies."""
    nvtx = nsys_nvtx_span_us(path, range_name)
    if nvtx:
        return nvtx + (f"{range_name} NVTX ranges",)
    return nsys_node_replay_us(path) + ("graph replays with copies",)


def iteration_note(measured):
    """Note for an end-to-end row from nsys_iteration's result per mode."""
    counts = " / ".join(str(measured[m][1]) for m in measured)
    modes = " (graph / direct)" if len(measured) > 1 else ""
    units = {measured[m][2] for m in measured}
    unit = units.pop() if len(units) == 1 else " / ".join(measured[m][2] for m in measured)
    return f"Median GPU span of the {counts}{modes} {unit} in the nsys traces."


def nsys_log_ms_per_iter(sqlite_path):
    """gpu_ms_per_iter the binary printed during an nsys run (the .log beside the trace)."""
    log = (REPO / sqlite_path).with_suffix(".log")
    return float(re.search(r"gpu_ms_per_iter=([\d.]+)", log.read_text()).group(1))


def loom_benchmark(path):
    runs = json.loads((REPO / path).read_text())
    return {(r["shape"], r["inputs"]): r["summary"] for r in runs}


def sweep_row(path, **match):
    with open(REPO / path) as f:
        for row in csv.DictReader(f):
            if all(row[k] == str(v) for k, v in match.items()):
                return row
    raise ValueError(f"{match} not in {path}")


# Per-run results in a time_all.sh log, by the line each harness prints:
# (regex, full batch scale to us, single scale to us). The groups are full
# batch, iterations in it, single iteration.
TIMING_LINES = [
    # llama2c: key=value fields
    (r"launches=(\d+) .*?gpu_ms=([\d.]+) .*?single_us=([\d.]+)", None, 1.0),
    # loom_msa: metrics: ... replays=N ... batch_ms=X single_us=Y
    (r"^metrics: .*?replays=(\d+) .*?batch_ms=([\d.]+) single_us=([\d.]+)", None, 1.0),
    # pagerank
    (r"^batch: ([\d.]+) ms for (\d+) back-to-back .*?single: ([\d.]+) ms", 1e3, 1e3),
    # mlperf: two lines
    (r"^  batch +([\d.]+) ms \((\d+) back-to-back[^\n]*\n  single +([\d.]+) us", 1e3, 1.0),
]


def timing_runs(path):
    """[(full batch us, iterations, single iteration us)] per run in a
    time_all.sh log."""
    text = (REPO / path).read_text()
    runs = []
    for run in text.split("### run ")[1:]:
        for pattern, batch_scale, single_scale in TIMING_LINES:
            m = re.search(pattern, run, re.M)
            if not m:
                continue
            if batch_scale is None:  # (iterations, batch ms, single)
                n, batch, single = int(m[1]), float(m[2]) * 1e3, float(m[3])
            else:
                batch, n, single = float(m[1]) * batch_scale, int(m[2]), float(m[3])
            runs.append((batch, n, single * single_scale))
            break
        else:
            raise ValueError(f"no batch / single timing in a run of {path}")
    if len({n for _, n, _ in runs}) != 1:
        raise ValueError(f"runs of {path} differ in batch size")
    return runs


# ------------------------------------------------------------------ cells ---

class Cell:
    """One measurement and the files behind it. Kernel cells hold one ncu time
    (us = single_us); whole-application cells hold the time per iteration in a
    batch (us) and, where time_all.sh ran, the full batch and the single-
    iteration time."""

    def __init__(self, us=None, ncu=None, nsys=None, source=None, empty="", predict=None,
                 kernel=False, batch_us=None, batch_n=None, single_us=None, runs=None):
        self.us, self.ncu, self.nsys, self.source = us, ncu, nsys, source
        self.empty, self.kernel = empty, kernel
        self.batch_us, self.batch_n, self.runs = batch_us, batch_n, runs
        self.single_us = us if kernel else single_us
        # Agentic optimized cells: the CUDA Graph version's cell (whole
        # application) or summed kernel time (kernel rows) for the same work.
        self.vs = None
        # How to predict this value on another GPU from this GPU's ncu report:
        # ("kernel", report, kernel, grid) or ("app", report, runs in report).
        self.predict = predict

    def value_file(self):
        """The file the value itself was read from."""
        return self.ncu or self.source or self.nsys

    def batched(self):
        """Per-iteration time in a timed batch, if this cell has one."""
        return self.us if self.batch_us is not None else None

    def overhead(self):
        """Single iteration / per iteration in the batch."""
        return ratio(self.single_us, self.batched())

    def agentic_speedups(self):
        """(full batch, single iteration) speedup over the CUDA Graph version."""
        if isinstance(self.vs, Cell):
            return ratio(self.vs.batched(), self.batched()), ratio(self.vs.single_us, self.single_us)
        return None, ratio(self.vs, self.single_us)


def ratio(a, b):
    return a / b if a is not None and b is not None else None


def graph_speedups(direct, graph):
    """(full batch, single iteration) speedup of the CUDA Graph over Direct."""
    return ratio(direct.batched(), graph.batched()), ratio(direct.single_us, graph.single_us)


NCU_PENDING = "pending (ncu not run yet)"
DIRECT_PENDING = "pending (direct mode not profiled yet)"
SOTA_PENDING = "pending (agentic optimized not profiled yet)"
TIMING_PENDING = "pending (time_all.sh not run yet)"
PENDING = (NCU_PENDING, DIRECT_PENDING, SOTA_PENDING, TIMING_PENDING)


def ncu_cell(path, kernel, grid=None, nsys=None, pending=NCU_PENDING):
    if path is None:
        return Cell(nsys=nsys, empty=pending, kernel=True)
    return Cell(ncu_kernel_us(path, kernel, grid)[0], ncu=path, nsys=nsys, kernel=True,
                predict=("kernel", path, kernel, grid))


def ncu_iteration_us(path, kernels, iterations=1, grid=None):
    """Summed ncu time of every launch of `kernels` (name prefixes) in a
    report covering `iterations` iterations, per iteration."""
    launches = [r for k in kernels for r in kernel_launches(path, k, grid)]
    if not launches:
        raise ValueError(f"{kernels} (grid={grid}) not found in {path}")
    return sum(launch_us(r) for r in launches) / iterations


def graph_kernels_us(cell, kernels):
    """CUDA Graph time of the baseline kernels an agentic optimized kernel row
    replaces: the sum of each kernel's mean launch time in the graph cell's
    report."""
    if cell.ncu is None:
        return None
    return sum(ncu_kernel_us(cell.ncu, k)[0] for k in kernels)


def sota_kernel_cell(gpu, key, kernels, vs, iterations=1, grid=None):
    path, nsys = gpu.file(f"{key}_sota_ncu"), gpu.file(f"{key}_sota_nsys")
    if path is None:
        cell = Cell(nsys=nsys, empty=SOTA_PENDING, kernel=True)
    else:
        cell = Cell(ncu_iteration_us(path, kernels, iterations, grid), ncu=path, nsys=nsys,
                    kernel=True)
    cell.vs = vs
    return cell


def timed_cell(log, nsys, predict=None):
    """Whole-application cell from a time_all.sh log: the medians over its
    runs of the full batch and the single-iteration time."""
    runs = timing_runs(log)
    batch = statistics.median(b for b, _, _ in runs)
    n = runs[0][1]
    return Cell(batch / n, source=log, nsys=nsys, predict=predict, batch_us=batch, batch_n=n,
                single_us=statistics.median(s for _, _, s in runs), runs=len(runs))


def app_cells(gpu, key, predict, fallback):
    """Whole-application cells per launch mode: from the GPU's time_all.sh
    logs where it has them, else fallback(mode) (a per-iteration time only)."""
    cells = {}
    for m in MODES:
        log, nsys = gpu.file(f"{key}_timing_{m}"), gpu.file(f"{key}_nsys_{m}")
        cells[m] = timed_cell(log, nsys, predict(m)) if log else fallback(m)
    return cells


def sota_app_cell(gpu, key, vs, fallback=None):
    """Agentic optimized whole-application cell: from its time_all.sh log, else
    fallback() if given, else pending."""
    log, nsys = gpu.file(f"{key}_sota_timing"), gpu.file(f"{key}_sota_nsys")
    if log:
        cell = timed_cell(log, nsys)
    elif fallback:
        cell = fallback()
    else:
        cell = Cell(nsys=nsys, empty=SOTA_PENDING)
    cell.vs = vs
    return cell


def timing_note(cell, what):
    """Note for a whole-application row timed by time_all.sh."""
    return (f"{what} Full batch: {cell.batch_n} back-to-back iterations between CUDA events "
            f"(per iteration = full batch / {cell.batch_n}). Single iteration: median of "
            f"{SINGLE_ITERATIONS} iterations, each launched alone on an idle device, so its "
            "launch cost is not hidden behind the previous iteration's work. Each time is the "
            f"median of {cell.runs} interleaved runs (time_all.sh).")


def sota_note(key, text):
    return f" Agentic optimized (launch={SOTA_LAUNCH[key]}): {text}"


# ------------------------------------------------------------------- rows ---
# Each row: (workload, kernel/stage, problem size, metric, direct, graph,
# agentic optimized, notes), the cells in column order. Per workload, at the
# largest problem size tested: the dominant kernel, when one kernel
# dominates, and the whole application; then the whole application at a small
# shape (small_row).

KERNEL_TIME = "Kernel time (ncu, one launch)"
WHOLE_APP = "Whole application"
TIMED = "Batch and single iteration (cudaEvent)"


def app_metric(cells, fallback):
    return TIMED if cells["graph"].batch_us is not None else fallback


def share(kernel, app, launches=1):
    if kernel.us is None or app.us is None:
        return ""
    return f"{100 * launches * kernel.us / app.us:.0f}% of the whole-application time per iteration."


def small_row(gpu, key, workload, size, what):
    """Whole-application row at time_all.sh's small shape for `key`, where
    launch cost dominates; pending on a GPU without those logs."""
    small = key + SMALL

    def pending(m):
        return Cell(nsys=gpu.file(f"{small}_nsys_{m}"), empty=TIMING_PENDING)
    app = app_cells(gpu, small, lambda m: None, pending)
    sota = sota_app_cell(gpu, small, app["graph"],
                         lambda: Cell(nsys=gpu.file(f"{small}_sota_nsys"), empty=TIMING_PENDING))
    why = ("Small shape: compare isolated latency with batched time in the latency sheet. "
           "Their ratio reflects scheduling and measurement effects, not pure launch overhead.")
    note = timing_note(app["graph"], what) if app["graph"].batch_us is not None else what
    return (workload, f"{WHOLE_APP}, small shape (launch-bound)", size, TIMED,
            app["direct"], app["graph"], sota,
            f"{why} {note}" + sota_note(key, f"{SMALL_TIMING_NAMES[small][2]}, the same "
                                        "implementation at this shape."))


def llama2c_rows(gpu):
    w = "llama2c (coalesced matmul)"
    size = "batch=1024, dim=16384, rows=8192"
    nsys = {m: gpu.file(f"llama_nsys_{m}") for m in MODES}
    ncu = {m: gpu.file(f"llama_ncu_{m}") for m in MODES}
    sweep = gpu.file("llama_sweep")

    def from_sweep(m):
        if not sweep:
            return Cell(empty=TIMING_PENDING)
        row = sweep_row(sweep, batch=1024, size=16384, rows=8192, repeats=1)
        # The ncu report covers 3 iterations.
        return Cell(float(row[f"{m}_gpu_ms"]) * 1e3, source=sweep, nsys=nsys[m],
                    predict=("app", ncu[m], 3))
    app = app_cells(gpu, "llama", lambda m: ("app", ncu[m], 3), from_sweep)
    matmul = {m: ncu_cell(ncu[m], "matmul_coalesced_kernel", nsys=nsys[m]) for m in MODES}
    if ncu["graph"]:
        s = ncu_time_shares(ncu["graph"])
        largest = (f"Largest kernel ({s['matmul_coalesced_kernel']:.0f}% of serialized kernel "
                   f"time under ncu; rmsnorm {s['rmsnorm_kernel']:.0f}%, softmax "
                   f"{s['softmax_kernel']:.0f}%). ")
    else:
        largest = "Largest kernel. "
    kernels = ("rmsnorm_kernel", "softmax_kernel", "matmul_coalesced_kernel")
    # The agentic optimized ncu report covers 3 iterations of one fused launch each.
    sota_matmul = sota_kernel_cell(gpu, "llama", ("llama2c_sota_kernel",),
                                   graph_kernels_us(matmul["graph"], kernels), iterations=3)
    runs = gpu.file("llama_sota_runs")

    def from_runs():
        times = [float(t) for t in re.findall(r"gpu_ms_per_iter=([\d.]+)", (REPO / runs).read_text())]
        return Cell(1e3 * statistics.mean(times), source=runs, nsys=gpu.file("llama_sota_nsys"))
    sota_app = sota_app_cell(gpu, "llama", app["graph"], from_runs if runs else None)
    if app["graph"].batch_us is not None:
        app_note = timing_note(app["graph"], "One iteration: rmsnorm + softmax + matmul on "
                               "device-resident data, no per-iteration copies.")
    elif sweep and all(nsys.values()):
        reruns = " / ".join(f"{1e3 * nsys_log_ms_per_iter(nsys[m]):.1f}" for m in MODES)
        app_note = ("Mean of 3 interleaved trials from the GPU-filling sweep; no per-iteration "
                    f"copies. The nsys traces rerun this configuration ({reruns} us per "
                    "iteration, graph / direct).")
    else:
        app_note = "No whole-application timing selected yet."
    return [
        (w, "matmul_coalesced_kernel", size, KERNEL_TIME, matmul["direct"], matmul["graph"],
         sota_matmul,
         largest + "It is a matrix-vector product, so its time does not grow with batch; the "
         "three kernels run concurrently on separate streams, so their times overlap. "
         + share(matmul["graph"], app["graph"])
         + sota_note("llama", "one fused kernel (llama2c_sota_kernel, one block per 64 KiB row, "
                     "each row read once from registers) computes rmsnorm, softmax and the "
                     "matmul, so its time is compared with the three graph kernels' summed ncu "
                     + (f"time ({sota_matmul.vs:,.1f} us)." if sota_matmul.vs else "time."))),
        (w, f"{WHOLE_APP} (rmsnorm + softmax + matmul iteration)", size,
         app_metric(app, "Iteration time (cudaEvent)"), app["direct"], app["graph"], sota_app,
         app_note + sota_note("llama", "llama2c_sota, the same device-resident nocopy "
                              "configuration, one fused kernel per iteration.")),
        small_row(gpu, "llama", w, "batch=64, dim=256, rows=512, 32 repeats",
                  "One iteration: 32 repeats of rmsnorm + softmax + matmul (96 kernels of about "
                  "a microsecond each) on device-resident data, no per-iteration copies."),
    ]


def loom_rows(gpu):
    w = "loom_msa (multi-sensor attention)"
    size = "25,000 queries x 250,000 context"
    nsys = {m: gpu.file(f"loom_nsys_{m}") for m in MODES}
    ncu = {m: gpu.file(f"loom_ncu_{m}") for m in MODES}
    bench = gpu.file("loom_benchmark")
    summary = loom_benchmark(bench)[("1x25000x250000", "random")] if bench else {}

    def from_benchmark(m):
        run = summary.get("optimized" if m == "graph" else f"optimized_{m}")
        return (Cell(run["end_to_end_us"]["median"], source=gpu.file("loom_benchmark_txt"),
                     nsys=nsys[m], predict=("app", ncu[m], 1)) if run
                else Cell(nsys=nsys[m], empty=DIRECT_PENDING))
    app = app_cells(gpu, "loom", lambda m: ("app", ncu[m], 1), from_benchmark)
    attention = {m: ncu_cell(ncu[m], "fuse_attention_kernel", nsys=nsys[m],
                             pending=NCU_PENDING if m == "graph" or nsys[m] else DIRECT_PENDING)
                 for m in MODES}
    kernels = ("project_audio_kernel", "project_imu_kernel", "fuse_attention_kernel",
               "reduce_statistics_kernel")
    sota_attention = sota_kernel_cell(gpu, "loom", ("msa_sota_attention_kernel",),
                                      graph_kernels_us(attention["graph"], kernels))
    launch = SOTA_LAUNCH["loom"]
    run = summary.get("sota" if launch == "graph" else f"sota_{launch}")
    sota_app = sota_app_cell(gpu, "loom", app["graph"], (lambda: Cell(
        run["end_to_end_us"]["median"], source=gpu.file("loom_benchmark_txt"),
        nsys=gpu.file("loom_sota_nsys"))) if run else None)
    if app["graph"].batch_us is not None:
        app_note = timing_note(app["graph"], "msa_optimized_large on random inputs, including "
                               "the H2D input and D2H output copies.")
    elif summary:
        modes = "both launch modes" if app["direct"].us is not None else "CUDA graph only"
        app_note = (f"benchmark.py: median of 5 runs, random inputs, {modes}. Includes H2D input "
                    "and D2H output copies.")
    else:
        app_note = "No whole-application timing selected yet."
    return [
        (w, "fuse_attention_kernel", size, KERNEL_TIME, attention["direct"], attention["graph"],
         sota_attention,
         "Single-pass online softmax, shared-memory context tiles, 8-way context split "
         "(FlashAttention-style). " + share(attention["graph"], app["graph"])
         + sota_note("loom", "msa_sota_attention_kernel (13-way context split, 8 queries per "
                     "thread) with the projections folded into its loads and the split merge and "
                     "statistics into its last block. msa_sota_large runs it as 13 concurrent "
                     "launches, one per imu copy chunk, which ncu would serialize, so ncu profiles "
                     "msa_sota_large_1chunk, the same kernel as one launch. Compared with the four "
                     "graph kernels' summed ncu "
                     + (f"time ({sota_attention.vs:,.1f} us)." if sota_attention.vs else "time."))),
        (w, f"{WHOLE_APP} (H2D + kernels + D2H)", size, app_metric(app, "Iteration time (cudaEvent)"),
         app["direct"], app["graph"], sota_app,
         app_note + sota_note("loom", "msa_sota_large, rotating over 3 random input sets; "
                              "output written straight to pinned host memory.")),
        small_row(gpu, "loom", w, "4 queries x 4 context (main.c)",
                  "msa_optimized_4x4 (main.c's shape) on random inputs, including the H2D input "
                  "and D2H output copies; every copy and kernel moves a few hundred bytes."),
    ]


def pagerank_rows(gpu):
    w = "pagerank (GAPBS)"
    size = "2^20 vertices, 16 edges/vertex, 20 iterations"
    ncu = {m: gpu.file(f"pagerank_ncu_{m}") for m in MODES}
    nsys = {m: gpu.file(f"pagerank_nsys_{m}") for m in MODES}
    measured = {}

    def from_nsys(m):
        if not nsys[m]:
            return Cell(empty=DIRECT_PENDING)
        measured[m] = nsys_iteration(nsys[m], "pagerank.full")
        return Cell(measured[m][0], nsys=nsys[m], predict=("app", ncu[m], 1))
    app = app_cells(gpu, "pagerank", lambda m: ("app", ncu[m], 1), from_nsys)
    update = {m: ncu_cell(ncu[m], "pagerank_update", grid=32768, nsys=nsys[m],
                          pending=NCU_PENDING if m == "graph" or nsys[m] else DIRECT_PENDING)
              for m in MODES}
    if ncu["graph"]:
        chunks = [launch_us(r) for r in kernel_launches(ncu["graph"], "pagerank_update")
                  if number(r["launch__grid_size"][0]) != 32768]
        first = (f"the first iteration runs as {len(chunks)} upload-overlapped chunks of "
                 f"~{statistics.mean(chunks):.0f} us each. 19 full iterations = ")
    else:
        first = "the first iteration runs as 8 upload-overlapped chunks. "
    if gpu.file("pagerank_sota_ncu"):
        # Iterations 1-19 (the first runs on the CSR during the upload); the
        # last of them runs as two launches, one wave over half the parts each.
        sota_update = sota_kernel_cell(gpu, "pagerank", ("binned_update_kernel",),
                                       update["graph"].us, iterations=19)
    else:
        sota_update = Cell(empty=SOTA_PENDING, kernel=True)
    sota_nsys = gpu.file("pagerank_sota_nsys")

    def sota_from_nsys():
        if not sota_nsys:
            return Cell(empty=SOTA_PENDING)
        return Cell(nsys_iteration(sota_nsys, "pagerank.full")[0], nsys=sota_nsys)
    sota_app = sota_app_cell(gpu, "pagerank", app["graph"], sota_from_nsys)
    what = ("One iteration is one PageRank run: the CSR graph uploaded from pinned memory in 8 "
            "chunks overlapped with the first of 20 iterations, then the scores downloaded.")
    if app["graph"].batch_us is not None:
        app_note = timing_note(app["graph"], what)
    elif measured:
        app_note = (iteration_note(measured) + " The CSR graph is uploaded from pinned memory in 8 "
                    "chunks overlapped with the first iteration. " + gpu.notes.get("pagerank_timing", ""))
    else:
        app_note = "No whole-application timing selected yet."
    return [
        (w, "pagerank_update (one iteration)", size, KERNEL_TIME, update["direct"], update["graph"],
         sota_update,
         "Pull gather of precomputed contributions. Full-graph iterations only; "
         + first + share(update["graph"], app["graph"], 19)
         + sota_note("pagerank", "binned_update_kernel, summed over iterations 1-19 / 19: a "
                     "cooperative pull over destination parts that fit in shared memory, with "
                     "edges sorted by source sector on the GPU during the upload (5.4M instead "
                     "of 16.8M gathered sectors per iteration).")),
        (w, f"{WHOLE_APP} (H2D + 20 iterations + D2H)", size,
         app_metric(app, "Iteration time (nsys)"), app["direct"], app["graph"], sota_app,
         app_note + sota_note("pagerank", "rotating over 3 inputs (the benchmark graph and two "
                              "random graphs of the same size); includes the on-GPU edge sort.")),
        small_row(gpu, "pagerank", w, "2^14 vertices, 16 edges/vertex, 20 iterations",
                  "The same PageRank run on a 2^14-vertex graph (make -C pagerank small): upload, "
                  "20 iterations of a few microseconds each, download."),
    ]


def mlperf_rows(gpu):
    rows = []
    for w, kernel, size, small_size, key, replaced, sota_kernel, kernel_text, app_text in (
            ("mlperf-tiny anomaly", "mse_partial_kernel", "64M elements", "16 elements (smoke)",
             "anomaly",
             ("mse_partial_kernel", "mse_finalize_kernel"), "mse_sota_kernel",
             "mse_sota_kernel, one single-pass reduction (the last block finishes the sum), "
             "compared with mse_partial_kernel + mse_finalize_kernel.",
             "the float input is copied in shrinking chunks while a second stream reduces each "
             "chunk as it lands; the result is written straight to pinned host memory."),
            ("mlperf-tiny KWS", "kws_postprocess_kernel", "12 classes x 4M rows",
             "12 classes x 1 row (smoke)", "kws",
             ("kws_postprocess_kernel",), "kws_postprocess_sota_kernel",
             "kws_postprocess_sota_kernel on device buffers (kernels-only), one row per thread "
             "with rows staged through shared memory for coalesced 16 B loads and stores.",
             "one zero-copy kernel reads the int8 rows from and writes the probabilities and "
             "argmax to mapped pinned host memory, so both PCIe directions stream at once."),
            ("mlperf-tiny VWW", "vww_preprocess_kernel", "256M bytes", "16 bytes (smoke)", "vww",
             ("vww_preprocess_kernel",), "vww_preprocess_sota_kernel",
             "vww_preprocess_sota_kernel on device buffers (kernels-only), four 16 B vectors "
             "per thread.",
             "64 chunks of 4 MiB on separate copy-in, compute and copy-out streams, so H2D and "
             "D2H overlap.")):
        ncu = {m: gpu.file(f"{key}_ncu_{m}") for m in MODES}
        nsys = {m: gpu.file(f"{key}_nsys_{m}") for m in MODES}
        measured = {}

        def from_nsys(m, key=key, ncu=ncu, nsys=nsys, measured=measured):
            if not nsys[m]:
                return Cell(empty=DIRECT_PENDING)
            measured[m] = nsys_iteration(nsys[m], f"{key}.end_to_end")
            return Cell(measured[m][0], nsys=nsys[m], predict=("app", ncu[m], 1))
        app = app_cells(gpu, key, lambda m, ncu=ncu: ("app", ncu[m], 1), from_nsys)
        k = {m: ncu_cell(ncu[m], kernel, nsys=nsys[m],
                         pending=NCU_PENDING if m == "graph" or nsys[m] else DIRECT_PENDING)
             for m in MODES}
        only = "" if kernel == "mse_partial_kernel" else "Only kernel in the application. "
        sota_k = sota_kernel_cell(gpu, key, (sota_kernel,), graph_kernels_us(k["graph"], replaced))
        sota_nsys = gpu.file(f"{key}_sota_nsys")

        def sota_from_nsys(key=key, sota_nsys=sota_nsys):
            if not sota_nsys:
                return Cell(empty=SOTA_PENDING)
            return Cell(nsys_iteration(sota_nsys, f"{key}.end_to_end")[0], nsys=sota_nsys)
        sota_app = sota_app_cell(gpu, key, app["graph"], sota_from_nsys)
        if app["graph"].batch_us is not None:
            app_note = timing_note(app["graph"], "Random inputs, including the H2D and D2H copies.")
        else:
            app_note = iteration_note(measured) if measured else "No whole-application timing selected yet."
        rows += [(w, kernel, size, KERNEL_TIME, k["direct"], k["graph"], sota_k,
                  only + share(k["graph"], app["graph"]) + sota_note(key, kernel_text)),
                 (w, f"{WHOLE_APP} (H2D + kernels + D2H)", size,
                  app_metric(app, "Iteration time (nsys)"), app["direct"], app["graph"], sota_app,
                  app_note + sota_note(key, "rotating over 3 inputs; " + app_text)),
                 small_row(gpu, key, w, small_size,
                           "The smoke size on random inputs, including the H2D and D2H copies "
                           "of a few bytes each.")]
    return rows


def gpu_rows(gpu):
    rows = llama2c_rows(gpu) + loom_rows(gpu) + pagerank_rows(gpu) + mlperf_rows(gpu)
    # Drop the agentic optimized part of the notes where it has no measurement yet.
    return [(*r[:-1], r[-1].split(" Agentic optimized (launch=")[0].rstrip()
             if r[6].us is None else r[-1]) for r in rows]


# ------------------------------------------------------------ predictions ---
# Predicted speedup of a GPU over BASELINE, from BASELINE's ncu reports and the
# two GPUs' published specs; METHOD explains it.

BOUND_PCT = 60  # a kernel below this on every unit is latency-bound

METHOD = [
    "The predicted speedup uses only the RTX 3090's ncu reports and the two GPUs' published "
    "specs (RTX 3090: 82 SMs, 1.695 GHz boost, 19.5 Gbps GDDR6X; RTX 4090: 128 SMs, 2.52 GHz "
    "boost, 21 Gbps GDDR6X; both 384-bit, PCIe 4.0 x16), so it is fixed before any 4090 "
    "measurement. Each unit's peak scales by its hardware ratio: SM compute by SMs x SM clock, "
    "L1/L2 cache throughput assumed to scale the same way, DRAM by memory data rate "
    "(the 3090's own SM and memory clocks come from its ncu report: 1.69 GHz and 18.98 Gbps, "
    "so the ratios are about 2.32x and 1.11x).",
    f"Kernel rows: ncu gives each launch's SM, memory (L1/L2) and DRAM throughput as a percent "
    f"of the 3090's peak. Keeping the work per unit fixed, a unit at u% that is r times faster "
    f"needs u/r% of the old time, and the slowest unit sets the new time: speedup = max(u) / "
    f"max(u / r). So a DRAM-bound kernel gets ~1.11x and a compute- or cache-bound one ~2.32x. "
    f"A launch below {BOUND_PCT}% on every unit is latency-bound and gets the SM clock ratio "
    f"(1.49x). A kernel with several launches combines them weighted by time.",
    "Whole-application rows: the kernels in that workload's ncu report (one run of the "
    "application) scale by their predicted speedups; the rest of the 3090 time (H2D/D2H copies "
    "over the same PCIe 4.0 link, launch gaps) is assumed unchanged. Where the kernels overlap "
    "(llama2c's three streams; their summed time exceeds the iteration time), the whole time "
    "scales with the kernels.",
    "Not modelled: cache capacity (the 4090's L2 is 72 MB vs 6 MB, so data that spilled to DRAM "
    "on the 3090 can stay in L2), kernels overlapping copies (pagerank's first iteration), "
    "clocks above rated boost, and compiler differences (CUDA 13.3 vs 12.5). Measured = "
    "3090 time / 4090 time for the CUDA Graph column; kernel rows read pending until the 4090 "
    "ncu reports exist.",
]


def prediction_method(gpu):
    if not gpu.specs:
        return ["Predictions are not enabled for this GPU. The existing RTX 4090 model "
                "has not been validated for RTX 5090's cache, memory bus, or PCIe link. "
                "Measured comparisons against RTX 3090 will populate when selected "
                "RTX 5090 results exist. Pending cells are not measurements."]
    return METHOD


def pct(row, metric):
    return number(row[metric][0])


def launch_speedup(row, target):
    """Predicted speedup on `target` of one launch in a BASELINE ncu report,
    and the reason."""
    base, new = BASELINE.specs, target.specs
    sm_ghz = number(row["gpc__cycles_elapsed.avg.per_second"][0])
    dram_gbps = 2 * number(row["dram__cycles_elapsed.avg.per_second"][0])  # GDDR6X: 2 bits/cycle
    compute = new["sms"] * new["boost_ghz"] / (base["sms"] * sm_ghz)
    units = {  # percent of the baseline's peak, and how much faster that peak is on target
        "SM": (pct(row, "sm__throughput.avg.pct_of_peak_sustained_elapsed"), compute),
        "L1/L2": (pct(row, "gpu__compute_memory_throughput.avg.pct_of_peak_sustained_elapsed"),
                  compute),
        "DRAM": (pct(row, "gpu__dram_throughput.avg.pct_of_peak_sustained_elapsed"),
                 new["dram_gbps"] / dram_gbps),
    }
    busiest = max(units, key=lambda k: units[k][0])
    peak = units[busiest][0]
    if peak < BOUND_PCT:
        clock = new["boost_ghz"] / sm_ghz
        return clock, (f"Latency-bound (busiest unit {busiest} {peak:.0f}% < {BOUND_PCT}%): "
                       f"SM clock {new['boost_ghz']} / {sm_ghz:.2f} GHz = {clock:.2f}x.")
    limit = max(units, key=lambda k: units[k][0] / units[k][1])
    u, r = units[limit]
    why = {
        "SM": f"SMs x clock {new['sms']} x {new['boost_ghz']} / ({base['sms']} x {sm_ghz:.2f}) "
              f"GHz = {r:.2f}x",
        "L1/L2": f"assumed to scale with SMs x clock, {r:.2f}x",
        "DRAM": f"data rate {new['dram_gbps']:g} / {dram_gbps:.2f} Gbps = {r:.2f}x",
    }[limit]
    others = ", ".join(f"{k} {v[0]:.0f}%" for k, v in units.items() if k != limit)
    text = f"{limit}-bound ({limit} {u:.0f}% busy; {others}): {why}"
    if u < peak:
        text += f"; {busiest} is busier but speeds up more, so {limit} limits"
    return peak / (u / r), text + "."


def predicted_speedup(cell, target):
    """(speedup, reason) of a BASELINE cell on `target`, or None."""
    if not target.specs or cell.predict is None or cell.predict[1] is None:
        return None
    kind, path, *rest = cell.predict
    launches = kernel_launches(path, *rest) if kind == "kernel" else ncu_rows(path)
    times = [launch_us(r) for r in launches]
    preds = [launch_speedup(r, target) for r in launches]
    kernels = sum(times)
    new_kernels = sum(t / s for t, (s, _) in zip(times, preds))
    if kind == "kernel":
        return kernels / new_kernels, max(zip(times, preds))[1][1]
    runs, = rest
    kernels, new_kernels = kernels / runs, new_kernels / runs
    k = kernels / new_kernels
    if kernels >= cell.us:
        return k, (f"Kernels overlap (their summed ncu time is {100 * kernels / cell.us:.0f}% of "
                   f"this time), so all of it scales with the kernels' {k:.2f}x.")
    other = cell.us - kernels
    return cell.us / (other + new_kernels), (
        f"Kernels are {100 * kernels / cell.us:.0f}% of the {BASELINE.sheet} time and predicted "
        f"{k:.2f}x faster; the other {100 * other / cell.us:.0f}% (copies, gaps) is unchanged.")


# ----------------------------------------------------------------- styling ---

BLUE = Font(color="0563C1", underline="single")
BOLD = Font(bold=True)
GREY = Font(italic=True, color="808080")
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
GROUP_FILLS = ["FCE4D6", "E2EFDA", "DDEBF7"]
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
# Column order of the cells in every row: Direct, CUDA Graph, Agentic optimized.
GROUPS = ["Direct launch (no graph)", "CUDA Graph", "Agentic optimized (formerly SOTA)"]
SHORT = ["Direct", "Graph", "Agentic"]
TIMES = ["Full batch (ms)", "Per iteration in batch (us)", "Single iteration (us)"]
NOT_APPLICABLE = "—"


def put_link(ws, row, col, path, text):
    c = ws.cell(row, col, text)
    c.hyperlink = link(path)
    c.font = BLUE
    return c


def time_values(c):
    """The cell's (full batch ms, per iteration us, single iteration us) as
    numbers, or the text to show instead; None = blank."""
    if c.kernel:  # one ncu time: a single launch in isolation
        return NOT_APPLICABLE, NOT_APPLICABLE, c.us if c.us is not None else c.empty
    if c.us is None:
        return None, c.empty, None
    if c.batch_us is None:
        return TIMING_PENDING, c.us, TIMING_PENDING
    return c.batch_us / 1e3, c.us, c.single_us


def short_note(row):
    """Short visible note; the full explanation stays in a cell comment."""
    if not row[-1]:
        return ""
    workload, stage, _, _, direct, graph, optimized, _ = row
    key = next(k for prefix, k in (
        ("llama2c", "llama"), ("loom_msa", "loom"), ("pagerank", "pagerank"),
        ("mlperf-tiny anomaly", "anomaly"), ("mlperf-tiny KWS", "kws"),
        ("mlperf-tiny VWW", "vww")) if workload.startswith(prefix))
    if graph.kernel:
        base = {
            "llama": "Concurrent baseline kernels; optimized speedup uses their summed NCU time.",
            "loom": "Split attention; optimized speedup uses all four baseline kernels; NCU uses one chunk.",
            "pagerank": "Full-graph updates; optimized time averages iterations 1–19.",
            "anomaly": "Optimized reduction replaces both baseline kernels.",
            "kws": "Kernel-only timing on device buffers; optimized rows use coalesced accesses.",
            "vww": "Kernel-only timing on device buffers; optimized kernel processes vectors.",
        }[key]
        if optimized.us is None:
            base = {
                "llama": "Concurrent baseline kernels; matrix-vector product.",
                "loom": "Online softmax with an eight-way context split.",
                "pagerank": "Full-graph updates; first update overlaps upload.",
                "anomaly": "Partial MSE reduction; final reduction is separate.",
                "kws": "Single kernel; device buffers.",
                "vww": "Single kernel; device buffers.",
            }[key]
        return base
    small = "small shape" in stage
    base = ({
        "llama": "Device-resident inputs; 32 repeats per iteration.",
        "loom": "Includes input and output transfers; 4 × 4 attention.",
        "pagerank": "Includes upload, 20 PageRank updates, and download.",
        "anomaly": "Includes input and output transfers; 16 elements.",
        "kws": "Includes input and output transfers; one 12-class row.",
        "vww": "Includes input and output transfers; 16 bytes.",
    } if small else {
        "llama": "Device-resident inputs; optimized version fuses all three operations.",
        "loom": "Includes transfers; optimized output goes to pinned host memory.",
        "pagerank": "Includes upload, 20 updates, and download; optimized run includes edge sorting.",
        "anomaly": "Includes transfers; optimized reduction overlaps chunked upload.",
        "kws": "Includes transfers; optimized kernel uses mapped host memory.",
        "vww": "Includes transfers; optimized version overlaps upload, compute, and download.",
    })[key]
    if optimized.us is None and not small:
        base = {
            "llama": "Device-resident inputs; mean of three sweep trials.",
            "loom": "Includes input and output transfers; benchmark median.",
            "pagerank": "Includes transfers; median iteration span from Nsight Systems.",
            "anomaly": "Includes transfers; median iteration span from Nsight Systems.",
            "kws": "Includes transfers; median iteration span from Nsight Systems.",
            "vww": "Includes transfers; median iteration span from Nsight Systems.",
        }[key]
    counts = {c.batch_n for c in (direct, graph, optimized) if c.batch_n is not None}
    runs = {c.runs for c in (direct, graph, optimized) if c.runs is not None}
    if len(counts) == 1:
        base += f" {counts.pop():,} iterations/batch."
    if len(runs) == 1:
        base += f" Median of {runs.pop()} runs."
    if all(c.us is None for c in (direct, graph, optimized)):
        base = "Measurements pending. " + base
    return base


def table_sheet(wb, title, headers, widths, description):
    ws = wb.create_sheet(title)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.cell(1, 1, title).font = Font(bold=True, size=14)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ws.cell(2, 1, description).alignment = WRAP
    ws.row_dimensions[2].height = 42
    for i, (header, width) in enumerate(zip(headers, widths), 1):
        c = ws.cell(4, i, header)
        c.font, c.fill, c.border = HEADER_FONT, HEADER_FILL, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[4].height = 42
    ws.freeze_panes = "C5"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.print_title_rows = "1:4"
    return ws


def finish_table(ws):
    ws.auto_filter.ref = f"A4:{get_column_letter(ws.max_column)}{ws.max_row}"
    for row in ws.iter_rows(min_row=5):
        for c in row:
            c.alignment = WRAP
            c.border = BOX
        ws.row_dimensions[row[0].row].height = 58


def write_value(ws, row, col, value, source=None, speedup=False):
    c = ws.cell(row, col)
    if isinstance(value, (int, float)):
        c.value = round(value, 3)
        c.number_format = '0.00"x"' if speedup else '#,##0.00'
        if source:
            c.hyperlink = link(source)
            c.font = BLUE
    elif value is not None:
        c.value = "pending" if value in PENDING else value
        c.font = GREY
        if value in PENDING:
            c.comment = Comment(value, "build_spreadsheet.py")
    return c


def write_note(ws, row, col, data):
    c = ws.cell(row, col, short_note(data))
    c.comment = Comment(data[-1], "build_spreadsheet.py", width=480, height=280)
    return c


def primary_speedups(cells):
    direct, graph, optimized = cells
    if graph.kernel:
        return graph_speedups(direct, graph)[1], optimized.agentic_speedups()[1]
    return ratio(direct.us, graph.us), ratio(graph.us, optimized.us)


def compact_stage(stage):
    return "Application (small)" if "small shape" in stage else stage.replace(WHOLE_APP, "Application")


def build_summary(wb, gpu, rows, base_rows=None, kernel=False):
    """Compact application summary, or separate kernel comparison table."""
    compare = base_rows is not None
    title = f"Kernels {gpu.sheet}" if kernel else gpu.sheet
    headers = ["Workload", "Stage", "Problem size", "Metric", "Direct (us)", "Graph (us)",
               "Optimized (us)", "Graph speedup\nDirect / Graph", "Optimized speedup\nGraph / Optimized"]
    if compare:
        headers += [f"GPU speedup\n{BASELINE.sheet.removeprefix('RTX ')} / "
                    f"{gpu.sheet.removeprefix('RTX ')}\n(Graph)"]
    headers += ["Notes"]
    widths = [24, 30, 30, 20, 15, 15, 15, 15, 17] + ([18] if compare else []) + [48]
    desc = (f"{gpu.label}; {gpu.toolkit}. " +
            ("NCU kernel durations; optimized speedup uses the baseline kernels replaced. " if kernel else
             "Application time per iteration in a back-to-back batch; lower is better. ") +
            "Times link to measurements. Notes have full explanations in comments.")
    ws = table_sheet(wb, title, headers, widths, desc)
    for col, target in enumerate((f"Latency {gpu.sheet}", f"Kernels {gpu.sheet}", "Sources", "GPU analysis"), 1):
        c = ws.cell(3, col, target)
        c.hyperlink = f"#'{target}'!A1"
        c.font = BLUE
    for i, data in enumerate(rows):
        workload, stage, size, metric, *cells, _ = data
        if cells[1].kernel != kernel:
            continue
        r = ws.max_row + 1
        metric = "NCU kernel" if kernel else ("CUDA events (batched)" if metric == TIMED else metric)
        for col, value in enumerate((workload, compact_stage(stage), size, metric), 1):
            ws.cell(r, col, value)
        ws.cell(r, 1).font = BOLD
        for col, cell, fill in zip(range(5, 8), cells, GROUP_FILLS):
            write_value(ws, r, col, cell.us if cell.us is not None else cell.empty, cell.value_file())
            ws.cell(r, col).fill = PatternFill("solid", fgColor=fill)
        for col, speedup in enumerate(primary_speedups(cells), 8):
            write_value(ws, r, col, speedup, speedup=True)
        if kernel and cells[2].vs is not None:
            ws.cell(r, 9).comment = Comment(
                "Baseline: summed NCU time of the kernels replaced by the optimized version. "
                + data[-1], "build_spreadsheet.py", width=480, height=280)
        if compare:
            base = base_rows[i][5]
            write_value(ws, r, 10, ratio(base.us, cells[1].us) or "pending", speedup=True)
        write_note(ws, r, len(headers), data)
    finish_table(ws)
    return ws


def build_latency(wb, gpu, rows):
    headers = ["Workload", "Stage", "Problem size", "Version", "Batch iterations", "Runs",
               "Full batch (ms)", "Batched (us/iteration)", "Single (us)", "Single / batched",
               "Graph speedup (single)", "Optimized speedup (single)", "Notes"]
    ws = table_sheet(wb, f"Latency {gpu.sheet}", headers,
                     [24, 24, 30, 16, 14, 9, 16, 17, 15, 15, 17, 18, 48],
                     f"{gpu.label}. Single = median of {SINGLE_ITERATIONS} isolated iterations. "
                     "Single/batched reflects scheduling and timing conditions, not pure launch overhead. "
                     "Runs and batch counts apply to each version; missing measurements stay pending.")
    for data in rows:
        workload, stage, size, _, *cells, _ = data
        if cells[1].kernel:
            continue
        for version, cell in zip(SHORT, cells):
            r = ws.max_row + 1
            for col, v in enumerate((workload, compact_stage(stage), size, version, cell.batch_n, cell.runs), 1):
                ws.cell(r, col, v)
            full, per, single = time_values(cell)
            for col, value in enumerate((full, per, single, cell.overhead()), 7):
                write_value(ws, r, col, value, cell.value_file() if col < 10 else None, speedup=col == 10)
            if version == "Graph":
                write_value(ws, r, 11, graph_speedups(cells[0], cells[1])[1], speedup=True)
            if version == "Agentic":
                write_value(ws, r, 12, cell.agentic_speedups()[1], speedup=True)
            write_note(ws, r, 13, data)
    finish_table(ws)


def build_sources(wb, results):
    ws = table_sheet(wb, "Sources",
                     ["GPU", "Workload", "Stage", "Problem size", "Version", "Measurement", "NCU report", "Nsight Systems", "Notes"],
                     [15, 24, 30, 30, 15, 44, 44, 44, 48],
                     "Source files for every comparison. Filenames and summary timings link to the original evidence.")
    for gpu, rows in results:
        for data in rows:
            workload, stage, size, _, *cells, _ = data
            for version, cell in zip(SHORT, cells):
                r = ws.max_row + 1
                for col, v in enumerate((gpu.sheet, workload, compact_stage(stage), size, version), 1):
                    ws.cell(r, col, v)
                for col, path in enumerate((cell.value_file(), cell.ncu, cell.nsys), 6):
                    if path:
                        put_link(ws, r, col, path, Path(path).name)
                write_note(ws, r, 9, data)
    finish_table(ws)


def build_gpu_analysis(wb, results, predictions):
    ws = table_sheet(wb, "GPU analysis",
                     ["GPU", "Workload", "Stage", "Problem size", "Predicted speedup", "Measured speedup", "Notes"],
                     [15, 24, 30, 30, 18, 18, 48],
                     f"Speedup = {BASELINE.sheet} / target GPU for Graph timings. "
                     "Predictions use hardware scaling; hover Notes for the calculation. Missing measurements stay pending.")
    base_rows = dict(results)[BASELINE]
    for gpu, rows in results:
        if gpu is BASELINE:
            continue
        for i, data in enumerate(rows):
            workload, stage, size, _, _, graph, _, _ = data
            prediction = predictions[gpu][i]
            r = ws.max_row + 1
            for col, v in enumerate((gpu.sheet, workload, compact_stage(stage), size), 1):
                ws.cell(r, col, v)
            write_value(ws, r, 5, prediction[0] if prediction else None, speedup=True)
            write_value(ws, r, 6, ratio(base_rows[i][5].us, graph.us) or "pending", speedup=True)
            note = ("Predictions not enabled for this GPU." if not gpu.specs else
                    "Hardware scaling estimate; see comment for assumptions." if prediction else
                    "No prediction available.")
            c = ws.cell(r, 7, note)
            if not gpu.specs:
                c.comment = Comment("\n\n".join(prediction_method(gpu)),
                                    "build_spreadsheet.py", width=540, height=200)
            elif prediction:
                c.comment = Comment(prediction[1] + "\n\n" + "\n\n".join(prediction_method(gpu)),
                                    "build_spreadsheet.py", width=540, height=400)
    finish_table(ws)

NCU_DETAIL_METRICS = [
    ("gpu__time_duration.sum", "Duration"),
    ("sm__throughput.avg.pct_of_peak_sustained_elapsed", "SM throughput %"),
    ("gpu__compute_memory_throughput.avg.pct_of_peak_sustained_elapsed", "Memory throughput %"),
    ("gpu__dram_throughput.avg.pct_of_peak_sustained_elapsed", "DRAM throughput %"),
    ("sm__warps_active.avg.pct_of_peak_sustained_active", "Achieved occupancy %"),
    ("sm__maximum_warps_per_active_cycle_pct", "Theoretical occupancy %"),
    ("launch__registers_per_thread", "Registers / thread"),
    ("launch__grid_size", "Grid"),
    ("launch__block_size", "Block"),
    ("launch__waves_per_multiprocessor", "Waves / SM"),
    ("dram__bytes_read.sum", "DRAM read"),
    ("dram__bytes_write.sum", "DRAM write"),
]


def build_ncu_details(wb, gpu, rows):
    """Every launch in every ncu report a summary sheet uses, with key metrics."""
    ws = wb.create_sheet(f"ncu metrics ({gpu.sheet})")
    reports = []
    for _, _, _, _, *cells, _ in rows:
        for c in cells:
            if c.ncu and c.ncu not in reports:
                reports.append(c.ncu)
    ws.append(["ncu report", "Launch #", "Kernel"] + [label for _, label in NCU_DETAIL_METRICS])
    if not reports:
        ws.append([NCU_PENDING])
        ws.cell(2, 1).font = GREY
    for path in reports:
        for i, launch in enumerate(ncu_rows(path)):
            values = []
            for metric, _ in NCU_DETAIL_METRICS:
                if metric not in launch:
                    values.append("")
                    continue
                v, unit = launch[metric]
                try:
                    v = number(v)
                except ValueError:
                    pass
                values.append(f"{v} {unit}".strip() if unit and unit != "%" else v)
            ws.append([path, i, launch["Kernel Name"][0].split("(")[0]] + values)
            put_link(ws, ws.max_row, 1, path, Path(path).name)
    for c in ws[1]:
        c.font, c.fill = HEADER_FONT, HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, w in enumerate([52, 9, 40] + [16] * len(NCU_DETAIL_METRICS), 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def write_markdown(results, predictions):
    """Compact application summaries; the workbook holds supporting detail."""
    def timing(cell):
        if cell.us is None:
            return "pending"
        return f"[{cell.us:,.2f}]({link(cell.value_file())})"

    def speedup(value):
        return f"{value:.2f}x" if value is not None else "pending"

    lines = [
        "# CUDA performance results", "",
        "Application time per iteration in a back-to-back batch, in microseconds; lower is better. "
        "Each timing links to its measurement. Inputs are selected explicitly in "
        "`results/inputs/<gpu>.json`; new files do not change selections automatically. "
        "Generated from existing results; no benchmarks are rerun.", "",
        "**Direct** enqueues the baseline operations each iteration. **Graph** replays the same "
        "operations in a CUDA Graph. **Optimized** is the agentic optimized implementation "
        "(`*_sota.cu`), which may change kernels, data movement, and launch mode. "
        "Graph speedup = Direct / Graph; optimized speedup = Graph / Optimized.", "",
        "The [spreadsheet](cuda_performance.xlsx) has compact application summaries plus "
        "supporting sheets for isolated latency, kernel comparisons, source files, GPU predictions, "
        "and full NCU metrics. Notes are brief; cell comments retain the full explanations. "
        "Single/batched timing ratios reflect scheduling and measurement conditions, not pure launch overhead.", "",
        "All application timings include input and output transfers except llama2c, which uses "
        "device-resident data. Batch counts and run counts appear in the notes or latency sheets. "
        "Pending means the measurement is unavailable.",
    ]
    base_rows = dict(results)[BASELINE]
    for gpu, rows in results:
        compare = gpu is not BASELINE
        headers = ["Workload", "Stage", "Problem size", "Metric", "Direct (us)", "Graph (us)",
                   "Optimized (us)", "Graph speedup", "Optimized speedup"]
        if compare:
            headers += [f"Measured {BASELINE.sheet} / {gpu.sheet} (Graph)"]
        headers += ["Notes"]
        align = ["---"] * 4 + ["--:"] * (len(headers) - 5) + ["---"]
        lines += ["", f"## {gpu.sheet}", "",
                  f"GPU: NVIDIA GeForce {gpu.label}. Binaries built with {gpu.toolkit}.", "",
                  "| " + " | ".join(headers) + " |", "|" + "|".join(align) + "|"]
        for i, data in enumerate(rows):
            workload, stage, size, metric, *cells, _ = data
            if cells[1].kernel:
                continue
            metric = "CUDA events (batched)" if metric == TIMED else metric
            values = [workload, compact_stage(stage), size, metric] + [timing(c) for c in cells]
            values += [speedup(v) for v in primary_speedups(cells)]
            if compare:
                values += [speedup(ratio(base_rows[i][5].us, cells[1].us))]
            values += [short_note(data)]
            lines.append("| " + " | ".join(values) + " |")
    lines += ["", "GPU predictions and their assumptions are in the workbook's **GPU analysis** sheet.", "",
              "Regenerate (needs `ncu` on PATH and `pip install openpyxl`):", "",
              "```bash", "python3 results/build_spreadsheet.py", "```", ""]
    MD_OUT.write_text("\n".join(lines))

def main():
    validate_inputs(GPUS)
    results = [(gpu, gpu_rows(gpu)) for gpu in GPUS]
    base_rows = dict(results)[BASELINE]
    predictions = {}
    for gpu, rows in results:
        if gpu is not BASELINE:
            assert [r[:2] for r in rows] == [r[:2] for r in base_rows], "rows must line up"
            # r[5]: the CUDA Graph cell
            predictions[gpu] = [predicted_speedup(r[5], gpu) for r in base_rows]
    wb = Workbook()
    wb.remove(wb.active)
    for gpu, rows in results:
        build_summary(wb, gpu, rows, base_rows if gpu is not BASELINE else None)
    for gpu, rows in results:
        build_latency(wb, gpu, rows)
        build_summary(wb, gpu, rows, base_rows if gpu is not BASELINE else None, kernel=True)
    build_sources(wb, results)
    build_gpu_analysis(wb, results, predictions)
    for gpu, rows in results:
        build_ncu_details(wb, gpu, rows)
    wb.save(OUT)
    write_markdown(results, predictions)
    print(f"wrote {OUT.relative_to(REPO)} and {MD_OUT.relative_to(REPO)} "
          f"({', '.join(f'{len(rows)} {gpu.sheet} rows' for gpu, rows in results)})")


if __name__ == "__main__":
    main()
